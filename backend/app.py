"""TVE GROUP leveranciersbeheer, durable storage and real Excel import/export."""
import base64
import csv
import difflib
import hashlib
import hmac
import io
import json
import math
import os
import re
import secrets
import threading
from backend.storage import connect, DB_PATH, IS_POSTGRES, IS_VERCEL, STORAGE_NAME
import time
import uuid
from copy import deepcopy
from datetime import datetime, timezone
from decimal import Decimal
from pathlib import Path

from fastapi import FastAPI, HTTPException, Request, UploadFile, File, Depends
from fastapi.responses import Response, FileResponse
from fastapi.staticfiles import StaticFiles
from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parent.parent
TOKEN_PATH = DB_PATH.with_suffix('.token')
COMPANIES = ['TVE', 'FORTY FIVE', 'L&R', 'BERGSLOT', 'VDR']
app = FastAPI(title='TVE GROUP leveranciers', docs_url=None, redoc_url=None)


def uid():
    return str(uuid.uuid4())


def now():
    return datetime.now(timezone.utc).isoformat()


def norm(value):
    return str(value or '').strip().casefold()


def js(value):
    return json.dumps(value, ensure_ascii=False, separators=(',', ':'))


_initialized=False
_initialize_lock=threading.RLock()


def db():
    if not _initialized:
        with _initialize_lock:
            if not _initialized:initialize()
    return connect()


def seed():
    w = load_workbook(ROOT / 'data/startgegevens.xlsx', data_only=True)
    rows = [r for r in list(w['Blad1'].values)[2:] if norm(r[0])]
    unique = {norm(r[0]) for r in rows}
    if len(rows) != 63 or len(unique) != 56:
        raise RuntimeError(f'Startbestand wijkt af: {len(rows)} regels, {len(unique)} leveranciers')
    names = [str(r[0]).strip() for r in w['Blad3'].values if norm(r[0])]
    for r in rows:
        if str(r[1]).strip() not in names:
            names.append(str(r[1]).strip())
    state = {'suppliers': [], 'categories': [{'id': uid(), 'name': n} for n in names],
             'companies': [{'id': uid(), 'name': n} for n in COMPANIES], 'fields': [], 'rows': [], 'trash': [],
             'seed': {'filename': 'LEVERANCIERS TVE GROUP - september 2026.xlsx', 'rows': 63, 'suppliers': 56, 'imported_at': now()}}
    for label in ['Contactpersoon', 'E-mail', 'Telefoon', 'Website', 'Opmerkingen']:
        state['fields'].append({'id': uid(), 'name': label, 'type': 'text', 'options': []})
    suppliers = {}
    categories = {c['name']: c['id'] for c in state['categories']}
    for source_line, row in enumerate(rows, 3):
        key = norm(row[0])
        if key not in suppliers:
            suppliers[key] = uid()
            state['suppliers'].append({'id': suppliers[key], 'name': str(row[0]).strip()})
        state['rows'].append({'id': uid(), 'supplier_id': suppliers[key], 'category_id': categories[str(row[1]).strip()],
                              'company_ids': [state['companies'][j]['id'] for j, col in enumerate([2, 4, 6, 8, 10]) if norm(row[col]) == 'x'],
                              'values': {}, 'source_line': source_line})
    return state


def initialize():
    global _initialized
    if not IS_POSTGRES and not IS_VERCEL:
        DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    with connect() as c:
        if IS_POSTGRES:
            c.execute('BEGIN IMMEDIATE')
            schema=(ROOT/'supabase/schema.sql').read_text()
            if c.execute("SELECT 1 FROM pg_namespace WHERE nspname='tve'").fetchone():
                # PostgreSQL checks CREATE on the database even for an existing
                # schema. The runtime role owns tve but cannot create new schemas.
                schema=schema.replace('CREATE SCHEMA IF NOT EXISTS tve;','')
            c.executescript(schema)
        else:
            c.execute('PRAGMA journal_mode=WAL')
            c.executescript('''
        CREATE TABLE IF NOT EXISTS store(id INTEGER PRIMARY KEY CHECK(id=1), revision INTEGER NOT NULL, payload TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS audit(id TEXT PRIMARY KEY, at TEXT NOT NULL, actor TEXT NOT NULL, label TEXT NOT NULL, before_data TEXT NOT NULL, after_data TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS users(id TEXT PRIMARY KEY, email TEXT UNIQUE NOT NULL, password TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS sessions(token TEXT PRIMARY KEY, user_id TEXT NOT NULL REFERENCES users(id), expires REAL NOT NULL);
        CREATE TABLE IF NOT EXISTS imports(id TEXT PRIMARY KEY, user_id TEXT NOT NULL, payload TEXT NOT NULL, created REAL NOT NULL);
        CREATE TABLE IF NOT EXISTS rate_limits(key TEXT PRIMARY KEY, count INTEGER NOT NULL, until REAL NOT NULL);
        CREATE TABLE IF NOT EXISTS settings(key TEXT PRIMARY KEY, value TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS user_access(user_id TEXT PRIMARY KEY REFERENCES users(id) ON DELETE CASCADE, role TEXT NOT NULL CHECK(role IN ('admin','company')), company_ids TEXT NOT NULL, active INTEGER NOT NULL DEFAULT 1, version INTEGER NOT NULL DEFAULT 1, display_name TEXT NOT NULL DEFAULT '');
        CREATE TABLE IF NOT EXISTS invitations(id TEXT PRIMARY KEY, email TEXT NOT NULL, display_name TEXT NOT NULL, role TEXT NOT NULL CHECK(role IN ('admin','company')), company_ids TEXT NOT NULL, token_hash TEXT UNIQUE NOT NULL, created_at TEXT NOT NULL, expires REAL NOT NULL, invited_by TEXT NOT NULL REFERENCES users(id) ON DELETE CASCADE, used_at TEXT, canceled_at TEXT);
        CREATE TABLE IF NOT EXISTS access_audit(id TEXT PRIMARY KEY, at TEXT NOT NULL, actor TEXT NOT NULL, action TEXT NOT NULL, details TEXT NOT NULL);
        ''')
        if not c.execute("SELECT 1 FROM settings WHERE key='access_migrated_v1'").fetchone():
            # Existing owner accounts and sessions retain administrator access.
            c.execute("INSERT OR IGNORE INTO user_access(user_id,role,company_ids,active,version,display_name) SELECT id,'admin','[]',1,1,'' FROM users WHERE 1=1")
            c.execute('INSERT OR IGNORE INTO settings VALUES(?,?)',('access_migrated_v1','1'))
        if not c.execute('SELECT 1 FROM store').fetchone():
            initial = seed()
            c.execute('INSERT INTO store VALUES(1,1,?)', (js(initial),))
            c.execute('INSERT INTO audit(id,at,actor,label,before_data,after_data) VALUES(?,?,?,?,?,?)', (uid(), now(), 'Systeem', 'Startimport: 63 regels / 56 leveranciers', js(initial), js(initial)))
        if not c.execute('SELECT 1 FROM users').fetchone():
            if IS_POSTGRES:
                code=os.environ.get('TVE_SETUP_CODE')
                if code and not c.execute("SELECT 1 FROM settings WHERE key='setup_code_hash'").fetchone():
                    c.execute('INSERT INTO settings VALUES(?,?)',('setup_code_hash',hashlib.sha256(code.encode()).hexdigest()))
            elif not TOKEN_PATH.exists():
                TOKEN_PATH.write_text(secrets.token_urlsafe(32))
                TOKEN_PATH.chmod(0o600)
    _initialized=True


@app.middleware('http')
async def protection(request, call_next):
    # A separate origin cannot submit authenticated modifications (CSRF).
    origin = request.headers.get('origin')
    if request.method not in ['GET', 'HEAD', 'OPTIONS'] and origin:
        from urllib.parse import urlparse
        if urlparse(origin).netloc != request.headers.get('host'):
            return Response(js({'detail': 'Verzoek vanaf een andere website geweigerd.'}), status_code=403, media_type='application/json')
    try:
        response = await call_next(request)
    except RuntimeError:
        return Response(js({'detail':'De databaseverbinding is nog niet beschikbaar. Controleer de Supabase-configuratie.'}),status_code=503,media_type='application/json')
    response.headers['X-Content-Type-Options'] = 'nosniff'
    response.headers['X-Frame-Options'] = 'DENY'
    response.headers['Referrer-Policy'] = 'no-referrer'
    response.headers['Content-Security-Policy'] = "default-src 'self'; script-src 'self'; style-src 'self' 'unsafe-inline'; img-src 'self' data:; connect-src 'self'; frame-ancestors 'none'; object-src 'none'; base-uri 'self'"
    if request.url.path.startswith('/api'):
        response.headers['Cache-Control'] = 'no-store'
    return response


def user(request: Request):
    token = request.cookies.get('tve_session', '')
    with db() as c:
        u = c.execute('SELECT u.id FROM sessions s JOIN users u ON u.id=s.user_id WHERE s.token=? AND s.expires>?',
                      (hashlib.sha256(token.encode()).hexdigest(), time.time())).fetchone()
        if u:return current_access(c,u['id'])
    if not u:
        raise HTTPException(401, 'Log in om de leveranciersgegevens te bekijken.')


def current_access(c,user_id):
    record=c.execute('SELECT u.id,u.email,a.role,a.company_ids,a.active,a.version,a.display_name FROM users u JOIN user_access a ON a.user_id=u.id WHERE u.id=?',(user_id,)).fetchone()
    if not record or not record['active']:raise HTTPException(401,'Uw toegang is niet actief. Neem contact op met de beheerder.')
    return {**dict(record),'company_ids':json.loads(record['company_ids'])}


def admin(u=Depends(user)):
    if u['role']!='admin':raise HTTPException(403,'Deze handeling is alleen beschikbaar voor beheerders.')
    return u


def check_admin(c,u):
    return admin(current_access(c,u['id']))


def visible_row(row,u):
    return any(company_id in row['company_ids'] or company_id in row.get('company_prices',{}) for company_id in u['company_ids'])


def scoped_state(state,u):
    if u['role']=='admin':return state
    state=deepcopy(state);allowed=set(u['company_ids'])
    state['rows']=[row for row in state['rows'] if visible_row(row,u)]
    for row in state['rows']:
        row['company_ids']=[key for key in row['company_ids'] if key in allowed]
        row['company_prices']={key:value for key,value in row.get('company_prices',{}).items() if key in allowed}
        row['price_entered_at']={key:value for key,value in row.get('price_entered_at',{}).items() if key in allowed}
        row.pop('updated_by',None)
        row.pop('price',None)
    supplier_ids={row['supplier_id'] for row in state['rows']}
    state['suppliers']=[supplier for supplier in state['suppliers'] if supplier['id'] in supplier_ids]
    for supplier in state['suppliers']:
        supplier['prices']={key:value for key,value in supplier.get('prices',{}).items() if key in allowed}
    category_ids={row['category_id'] for row in state['rows']}
    state['categories']=[category for category in state['categories'] if category['id'] in category_ids]
    state['companies']=[company for company in state['companies'] if company['id'] in allowed]
    state['price_history']=[h for h in state.get('price_history',[]) if h['company_id'] in allowed]
    for row in state['rows']:
        row['price_effective_dates']={key:value for key,value in row.get('price_effective_dates',{}).items() if key in allowed}
    state['trash']=[];state.pop('seed',None)
    return state


def password_hash(password, salt=None):
    salt = salt or secrets.token_hex(16)
    digest = hashlib.scrypt(password.encode(), salt=bytes.fromhex(salt), n=16384, r=8, p=1).hex()
    return salt + ':' + digest


def session(response, request, user_id):
    token = secrets.token_urlsafe(40)
    with db() as c:
        c.execute('DELETE FROM sessions WHERE expires<?', (time.time(),))
        c.execute('INSERT INTO sessions VALUES(?,?,?)', (hashlib.sha256(token.encode()).hexdigest(), user_id, time.time()+43200))
    response.set_cookie('tve_session', token, httponly=True, samesite='lax', secure=request.headers.get('x-forwarded-proto') == 'https' or request.url.scheme == 'https', max_age=43200)


@app.get('/api/health')
def health():
    with db() as c:
        c.execute('SELECT revision FROM store').fetchone()
    return {'ok': True, 'storage': STORAGE_NAME}


@app.get('/api/auth/status')
def auth_status(request: Request):
    with db() as c:
        setup = not bool(c.execute('SELECT 1 FROM users').fetchone())
    try:
        u = user(request)
    except HTTPException:
        u = None
    return {'setup_required': setup, 'user': u}


@app.post('/api/auth/setup')
def setup_account(body: dict, response: Response, request: Request):
    email = str(body.get('email', '')).strip().lower()
    password = str(body.get('password', ''))
    if '@' not in email or len(password) < 12:
        raise HTTPException(422, 'Gebruik een geldig e-mailadres en een wachtwoord van minstens 12 tekens.')
    with db() as c:
        c.execute('BEGIN IMMEDIATE')
        if c.execute('SELECT 1 FROM users').fetchone():
            raise HTTPException(409, 'Het beheerdersaccount is al ingesteld. Log in.')
        supplied=str(body.get('code',''))
        stored=c.execute("SELECT value FROM settings WHERE key='setup_code_hash'").fetchone() if IS_POSTGRES else None
        valid=hmac.compare_digest(hashlib.sha256(supplied.encode()).hexdigest(),stored['value']) if stored else (not IS_POSTGRES and TOKEN_PATH.exists() and hmac.compare_digest(supplied,TOKEN_PATH.read_text().strip()))
        if not valid:
            raise HTTPException(403, 'Gebruik de privé-openingslink om het beheerdersaccount aan te maken.')
        user_id = uid()
        c.execute('INSERT INTO users VALUES(?,?,?)', (user_id, email, password_hash(password)))
        c.execute('INSERT INTO user_access VALUES(?,?,?,?,?,?)',(user_id,'admin','[]',1,1,''))
        c.execute("DELETE FROM settings WHERE key='setup_code_hash'")
    if not IS_POSTGRES:TOKEN_PATH.unlink(missing_ok=True)
    session(response, request, user_id)
    return {'ok': True}


@app.post('/api/auth/login')
def login(body: dict, response: Response, request: Request):
    email = str(body.get('email', '')).strip().lower()
    rate_key = hashlib.sha256(email.encode()).hexdigest()
    with db() as c:
        c.execute('BEGIN IMMEDIATE')
        limit = c.execute('SELECT * FROM rate_limits WHERE key=?', (rate_key,)).fetchone()
        if limit and limit['until'] > time.time() and limit['count'] >= 10:
            raise HTTPException(429, 'Te veel pogingen. Probeer het over 15 minuten opnieuw.')
        record = c.execute('SELECT * FROM users WHERE email=?', (email,)).fetchone()
        salt = record['password'].split(':')[0] if record else '00'*16
        candidate = password_hash(str(body.get('password', '')), salt)
        access=c.execute('SELECT active FROM user_access WHERE user_id=?',(record['id'],)).fetchone() if record else None
        valid = record and access and access['active'] and hmac.compare_digest(record['password'], candidate)
        if not valid:
            count = limit['count'] + 1 if limit and limit['until'] > time.time() else 1
            until = limit['until'] if limit and limit['until'] > time.time() else time.time()+900
            c.execute('INSERT OR REPLACE INTO rate_limits VALUES(?,?,?)', (rate_key, count, until))
        else:
            c.execute('DELETE FROM rate_limits WHERE key=?', (rate_key,))
    if not valid:
        raise HTTPException(401, 'E-mailadres of wachtwoord is onjuist.')
    session(response, request, record['id'])
    return {'ok': True}


@app.post('/api/auth/logout')
def logout(request: Request, response: Response):
    token = request.cookies.get('tve_session', '')
    with db() as c:
        c.execute('DELETE FROM sessions WHERE token=?', (hashlib.sha256(token.encode()).hexdigest(),))
    response.delete_cookie('tve_session')
    return {'ok': True}


def access_log(c,u,action,details):
    c.execute('INSERT INTO access_audit VALUES(?,?,?,?,?)',(uid(),now(),u['email'],action,js(details)))


def access_choice(state,body):
    role=body.get('role','company');company_ids=body.get('company_ids',[])
    if role not in ['admin','company'] or not isinstance(company_ids,list) or any(not isinstance(key,str) for key in company_ids):
        raise HTTPException(422,'Kies een geldige rol en bedrijven.')
    company_ids=list(dict.fromkeys(company_ids))
    if role=='admin':company_ids=[]
    elif not company_ids:raise HTTPException(422,'Kies minstens één bedrijf voor deze gebruiker.')
    for key in company_ids:entity(state,'companies',key)
    return role,company_ids


def invitation_status(record):
    return 'accepted' if record['used_at'] else 'canceled' if record['canceled_at'] else 'expired' if record['expires']<=time.time() else 'pending'


def valid_invitation(c,token):
    if not isinstance(token,str) or not 32<=len(token)<=128:raise HTTPException(410,'Deze uitnodigingslink is niet meer geldig. Vraag de beheerder om een nieuwe link.')
    record=c.execute('SELECT * FROM invitations WHERE token_hash=?',(hashlib.sha256(token.encode()).hexdigest(),)).fetchone()
    if not record or invitation_status(record)!='pending':raise HTTPException(410,'Deze uitnodigingslink is niet meer geldig. Vraag de beheerder om een nieuwe link.')
    return record


@app.get('/api/users')
def list_users(u=Depends(admin)):
    with db() as c:
        check_admin(c,u)
        users=c.execute('SELECT u.id,u.email,a.role,a.company_ids,a.active,a.version,a.display_name FROM users u JOIN user_access a ON a.user_id=u.id ORDER BY u.email').fetchall()
        invitations=c.execute('SELECT id,email,display_name,role,company_ids,created_at,expires,used_at,canceled_at FROM invitations ORDER BY created_at DESC').fetchall()
        audit=c.execute('SELECT at,actor,action,details FROM access_audit ORDER BY at DESC LIMIT 100').fetchall()
    return {'users':[{**dict(row),'company_ids':json.loads(row['company_ids'])} for row in users],
            'invitations':[{**dict(row),'company_ids':json.loads(row['company_ids']),'status':invitation_status(row)} for row in invitations],
            'history':[{**dict(row),'details':json.loads(row['details'])} for row in audit]}


@app.post('/api/users/invite')
def invite_user(body:dict,u=Depends(admin)):
    email=str(body.get('email','')).strip().lower();name=str(body.get('display_name','')).strip()
    if len(email)>250 or not re.fullmatch(r'[^\s@]+@[^\s@]+\.[^\s@]+',email):raise HTTPException(422,'Gebruik een geldig e-mailadres.')
    if len(name)>100:raise HTTPException(422,'Een naam mag maximaal 100 tekens bevatten.')
    token=secrets.token_urlsafe(32);invitation_id=uid();expires=time.time()+7*86400
    with db() as c:
        c.execute('BEGIN IMMEDIATE');u=check_admin(c,u)
        state,_=read_state(c);role,company_ids=access_choice(state,body)
        if c.execute('SELECT 1 FROM users WHERE email=?',(email,)).fetchone():raise HTTPException(409,'Dit e-mailadres heeft al een account. Pas diens toegang aan in de gebruikerslijst.')
        pending=c.execute('SELECT id FROM invitations WHERE email=? AND used_at IS NULL AND canceled_at IS NULL AND expires>? ORDER BY created_at DESC',(email,time.time())).fetchone()
        if pending and body.get('replace_id')!=pending['id']:
            raise HTTPException(409,'Dit e-mailadres heeft al een openstaande uitnodiging. Kies Nieuwe link maken bij die uitnodiging om haar bewust te vervangen.')
        if body.get('replace_id') and (not pending or pending['id']!=body['replace_id']):
            raise HTTPException(409,'De uitnodiging is intussen gewijzigd. Vernieuw de lijst voordat u een nieuwe link maakt.')
        c.execute('UPDATE invitations SET canceled_at=? WHERE email=? AND used_at IS NULL AND canceled_at IS NULL',(now(),email))
        c.execute('INSERT INTO invitations VALUES(?,?,?,?,?,?,?,?,?,?,?)',(invitation_id,email,name,role,js(company_ids),hashlib.sha256(token.encode()).hexdigest(),now(),expires,u['id'],None,None))
        access_log(c,u,'Uitnodiging gemaakt',{'email':email,'role':role,'company_ids':company_ids})
    # Only this response contains the secret; lists and the database store its hash.
    return {'id':invitation_id,'token':token,'email':email,'expires':expires}


@app.post('/api/users/invitations/{invitation_id}/cancel')
def cancel_invitation(invitation_id:str,u=Depends(admin)):
    with db() as c:
        c.execute('BEGIN IMMEDIATE');u=check_admin(c,u)
        record=c.execute('SELECT * FROM invitations WHERE id=?',(invitation_id,)).fetchone()
        if not record:raise HTTPException(404,'Uitnodiging niet gevonden.')
        if record['used_at']:raise HTTPException(422,'Deze uitnodiging is al geaccepteerd. U kunt het account deactiveren in de gebruikerslijst.')
        c.execute('UPDATE invitations SET canceled_at=? WHERE id=?',(now(),invitation_id))
        access_log(c,u,'Uitnodiging ingetrokken',{'email':record['email']})
    return {'ok':True}


@app.post('/api/users/{user_id}/access')
def change_access(user_id:str,body:dict,u=Depends(admin)):
    with db() as c:
        c.execute('BEGIN IMMEDIATE');u=check_admin(c,u)
        old=c.execute('SELECT u.email,a.* FROM users u JOIN user_access a ON a.user_id=u.id WHERE u.id=?',(user_id,)).fetchone()
        if not old:raise HTTPException(404,'Gebruiker niet gevonden.')
        if body.get('version')!=old['version']:raise HTTPException(409,'De toegang is intussen gewijzigd. Sluit dit formulier en open de gebruiker opnieuw om de nieuwste rechten te controleren.')
        state,_=read_state(c);role,company_ids=access_choice(state,body)
        active=body.get('active',True)
        if not isinstance(active,bool):raise HTTPException(422,'Kies een geldige toegangsstatus.')
        if user_id==u['id'] and (role!='admin' or not active):raise HTTPException(422,'U kunt uw eigen beheerderstoegang niet uitschakelen.')
        c.execute('UPDATE user_access SET role=?,company_ids=?,active=?,version=version+1 WHERE user_id=?',(role,js(company_ids),int(active),user_id))
        # Revoke every old session so changed roles take effect even in open tabs.
        if user_id!=u['id']:c.execute('DELETE FROM sessions WHERE user_id=?',(user_id,))
        access_log(c,u,'Toegang gewijzigd',{'email':old['email'],'role':role,'company_ids':company_ids,'active':active})
    return {'ok':True}


@app.post('/api/auth/invitation')
def inspect_invitation(body:dict):
    with db() as c:
        record=valid_invitation(c,body.get('token'));state,_=read_state(c)
        role,company_ids=access_choice(state,{'role':record['role'],'company_ids':json.loads(record['company_ids'])})
        companies=[entity(state,'companies',key)['name'] for key in company_ids]
    return {'email':record['email'],'display_name':record['display_name'],'role':role,'companies':companies,'expires':record['expires']}


@app.post('/api/auth/accept-invitation')
def accept_invitation(body:dict,response:Response,request:Request):
    password=body.get('password')
    if not isinstance(password,str) or not 12<=len(password)<=1024:raise HTTPException(422,'Kies een wachtwoord van minstens 12 en maximaal 1024 tekens.')
    with db() as c:
        c.execute('BEGIN IMMEDIATE');record=valid_invitation(c,body.get('token'));state,_=read_state(c)
        role,company_ids=access_choice(state,{'role':record['role'],'company_ids':json.loads(record['company_ids'])})
        if c.execute('SELECT 1 FROM users WHERE email=?',(record['email'],)).fetchone():raise HTTPException(409,'Dit e-mailadres heeft al een account. Log in met uw eigen wachtwoord.')
        user_id=uid()
        c.execute('INSERT INTO users VALUES(?,?,?)',(user_id,record['email'],password_hash(password)))
        c.execute('INSERT INTO user_access VALUES(?,?,?,?,?,?)',(user_id,role,js(company_ids),1,1,record['display_name']))
        c.execute('UPDATE invitations SET used_at=? WHERE id=?',(now(),record['id']))
        access_log(c,{'email':record['email']},'Uitnodiging geaccepteerd',{'role':role,'company_ids':company_ids})
    session(response,request,user_id)
    return {'ok':True}


def read_state(c):
    record = c.execute('SELECT * FROM store WHERE id=1').fetchone()
    return json.loads(record['payload']), record['revision']


@app.get('/api/state')
def get_state(u=Depends(user)):
    with db() as c:
        u=current_access(c,u['id'])
        state, revision = read_state(c)
    return {'state': scoped_state(state,u), 'revision': revision, 'user':u}


def entity(state, kind, entity_id):
    result = next((x for x in state[kind] if x['id'] == entity_id), None)
    if not result:
        raise HTTPException(422, f'Onbekende koppeling in {kind}. Vernieuw de gegevens.')
    return result


def price_amount(value):
    """Money is stored as decimal text, never as a binary floating point value."""
    if value is None or value == '':
        return ''
    if isinstance(value, bool) or not isinstance(value, (str, int, float)):
        raise HTTPException(422, 'Voer een geldig bedrag in euro in.')
    text = str(value).strip().replace('€', '').replace(' ', '').replace('\u00a0', '')
    if not text:
        return ''
    if ',' in text:
        if '.' in text and not re.fullmatch(r'\d{1,3}(?:\.\d{3})+,\d{1,2}', text):
            raise HTTPException(422, 'Gebruik bijvoorbeeld 1234,56 of 1.234,56 voor een prijs.')
        text = text.replace('.', '').replace(',', '.')
    if len(text) > 32 or not re.fullmatch(r'\d+(?:\.\d{1,2})?', text):
        raise HTTPException(422, 'Een prijs moet nul of hoger zijn, met maximaal twee decimalen.')
    amount = Decimal(text)
    if amount > Decimal('999999999999.99'):
        raise HTTPException(422, 'De prijs is te hoog (maximaal 999.999.999.999,99 euro).')
    return format(amount, '.2f')


def validate_prices(state):
    for supplier in state['suppliers']:
        prices = supplier.get('prices', {})
        if not isinstance(prices, dict):
            raise HTTPException(422, 'Ongeldige totaalprijzen bij een leverancier.')
        for company_id, amount in prices.items():
            entity(state, 'companies', company_id)
            price_amount(amount)


def update_product(row, values):
    if 'product_name' in values:
        name=values['product_name']
        if not isinstance(name,str) or len(name.strip())>250:
            raise HTTPException(422,'Een productnaam mag maximaal 250 tekens bevatten.')
        if name.strip():row['product_name']=name.strip()
        else:row.pop('product_name',None)
    if 'price' in values:
        amount=price_amount(values['price'])
        if amount!='':row['price']=amount
        else:row.pop('price',None)


def update_company_prices(state, row, values):
    if not isinstance(values, dict):
        raise HTTPException(422, 'Ongeldige productprijzen per bedrijf.')
    for company_id, value in values.items():
        entity(state, 'companies', company_id)
        amount = price_amount(value)
        if amount != '':row.setdefault('company_prices', {})[company_id] = amount
        elif 'company_prices' in row:row['company_prices'].pop(company_id, None)
    if not row.get('company_prices'):row.pop('company_prices', None)


def validate_state(state):
    required = ['suppliers', 'categories', 'companies', 'fields', 'rows', 'trash']
    if not isinstance(state, dict) or any(not isinstance(state.get(k), list) for k in required):
        raise HTTPException(422, 'Onjuist back-upformaat.')
    history=state.get('price_history',[])
    if not isinstance(history,list):raise HTTPException(422,'Ongeldige prijshistorie.')
    from backend.price_history import effective_date
    for h in history:
        if not isinstance(h,dict) or not all(isinstance(h.get(k),str) for k in ['row_id','company_id','date','amount','recorded_at']):
            raise HTTPException(422,'Ongeldige prijshistorie.')
        effective_date(h['date']);h['amount']=price_amount(h['amount'])
        try:datetime.fromisoformat(h['recorded_at'])
        except ValueError:raise HTTPException(422,'Ongeldige opslagdatum.')
    for kind in required:
        ids = [x.get('id') for x in state[kind] if isinstance(x, dict)]
        if len(ids) != len(state[kind]) or any(not isinstance(i,str) or not i for i in ids) or len(ids) != len(set(ids)):
            raise HTTPException(422, f'Ongeldige of dubbele ID’s in {kind}.')
    for kind in ['suppliers', 'categories', 'companies', 'fields']:
        for x in state[kind]:
            if not isinstance(x.get('name'),str) or not norm(x['name']) or len(x['name']) > 250:
                raise HTTPException(422, 'Een naam is verplicht (maximaal 250 tekens).')
    for field in state['fields']:
        if field.get('type') not in ['text','number','date','select'] or not isinstance(field.get('options'),list):
            raise HTTPException(422,'Ongeldig veldtype of keuzelijst.')
    for kind in ['categories','companies','fields']:
        names=[norm(x['name']) for x in state[kind]]
        if len(names)!=len(set(names)):
            raise HTTPException(422,'Deze naam bestaat al. Kies een andere naam.')
    validate_prices(state)
    for row in state['rows']:
        if 'product_name' in row and (not isinstance(row['product_name'],str) or len(row['product_name'])>250):
            raise HTTPException(422,'Ongeldige productnaam (maximaal 250 tekens).')
        if 'price' in row:price_amount(row['price'])
        if not isinstance(row.get('company_prices', {}), dict):
            raise HTTPException(422, 'Ongeldige productprijzen per bedrijf.')
        for company_id, amount in row.get('company_prices', {}).items():
            entity(state, 'companies', company_id)
            price_amount(amount)
        entity(state,'suppliers',row.get('supplier_id'))
        entity(state,'categories',row.get('category_id'))
        if not isinstance(row.get('company_ids'),list) or len(row['company_ids'])!=len(set(row['company_ids'])):
            raise HTTPException(422,'Ongeldige bedrijfskoppelingen.')
        for company in row['company_ids']:
            entity(state,'companies',company)
        if not isinstance(row.get('values'),dict):
            raise HTTPException(422,'Ongeldige extra velden.')
        for field_id,value in row['values'].items():
            f=entity(state,'fields',field_id)
            if value is None or value=='':
                continue
            if f['type']=='number':
                try:
                    if not math.isfinite(float(str(value).replace(',','.'))):
                        raise ValueError()
                except (ValueError,TypeError):
                    raise HTTPException(422,f'{f["name"]} moet een getal zijn.')
            elif f['type']=='date':
                try:
                    datetime.strptime(str(value),'%Y-%m-%d')
                except ValueError:
                    raise HTTPException(422,f'{f["name"]} moet een geldige datum zijn.')
            elif f['type']=='select' and value not in f['options']:
                raise HTTPException(422,f'{f["name"]}: kies een waarde uit de keuzelijst.')
            elif not isinstance(value,(str,int,float)):
                raise HTTPException(422,'Ongeldige veldwaarde.')


def save(c, before, state, revision, actor, label, record_dates=True):
    validate_prices(state)
    for supplier in state['suppliers']:
        if 'prices' in supplier:
            supplier['prices'] = {key: price_amount(value) for key, value in supplier['prices'].items() if price_amount(value) != ''}
    for row in state['rows']:
        update_product(row,{key:row[key] for key in ['product_name','price'] if key in row})
        update_company_prices(state, row, deepcopy(row.get('company_prices', {})))
        if isinstance(row.get('values'), dict):
            row['values'] = {key: value for key, value in row['values'].items() if value is not None and value != ''}
    if record_dates:
        timestamp = now()
        old_rows = {row['id']: row for row in before['rows']}
        for row in state['rows']:
            old = old_rows.get(row['id'])
            if old is None:
                row['created_at'] = timestamp
            if old != row:
                row['updated_at'] = timestamp
                row['updated_by'] = actor
            if row.get('price') is not None and (old is None or old.get('price') != row.get('price')):
                row.setdefault('price_entered_at', {})['general'] = timestamp
            for company_id, amount in row.get('company_prices', {}).items():
                if old is None or old.get('company_prices', {}).get(company_id) != amount:
                    row.setdefault('price_entered_at', {})[company_id] = timestamp
    if record_dates and before.get('price_history'):
        from backend.price_history import record_prices, ZONE
        today=datetime.now(ZONE).date().isoformat()
        for row in state['rows']:
            old=next((r for r in before['rows'] if r['id']==row['id']),{})
            pairs={h['company_id'] for h in before['price_history'] if h['row_id']==row['id']}
            for co in pairs:
                previous=old.get('price') if co=='general' else old.get('company_prices',{}).get(co)
                current=row.get('price') if co=='general' else row.get('company_prices',{}).get(co)
                old_history=[h for h in before['price_history'] if h['row_id']==row['id'] and h['company_id']==co]
                new_history=[h for h in state.get('price_history',[]) if h['row_id']==row['id'] and h['company_id']==co]
                if previous!=current and old_history==new_history and (co=='general' or co in {c['id'] for c in state['companies']}):
                    record_prices(state,before,row,{co:current or ''},today,now(),price_amount)
    validate_state(state)
    company_ids={company['id'] for company in state['companies']}
    for access in c.execute("SELECT company_ids FROM user_access WHERE role='company' AND active=1"):
        if not set(json.loads(access['company_ids']))<=company_ids:
            raise HTTPException(422,'Er zijn gebruikers met toegang tot een bedrijf dat ontbreekt. Wijs hun toegang eerst opnieuw toe via Gebruikers.')
    for invitation in c.execute("SELECT company_ids FROM invitations WHERE role='company' AND used_at IS NULL AND canceled_at IS NULL AND expires>?",(time.time(),)):
        if not set(json.loads(invitation['company_ids']))<=company_ids:
            raise HTTPException(422,'Er zijn actieve uitnodigingen voor een bedrijf dat ontbreekt. Trek die eerst in via Gebruikers of wijs het bedrijf opnieuw toe.')
    c.execute('UPDATE store SET revision=?,payload=? WHERE id=1', (revision+1, js(state)))
    c.execute('INSERT INTO audit(id,at,actor,label,before_data,after_data) VALUES(?,?,?,?,?,?)', (uid(), now(), actor, label, js(before), js(state)))
    return {'state':state,'revision':revision+1,'message':label}


def check_revision(body, revision):
    if body.get('revision') != revision:
        raise HTTPException(409, 'Iemand heeft de gegevens intussen gewijzigd. Uw invoer is niet overschreven. Vernieuw en controleer de nieuwste versie voordat u opnieuw opslaat.')


def find_or_create(state, kind, name):
    name = str(name or '').strip()
    if not name:
        raise HTTPException(422, 'Leverancier en productcategorie zijn verplicht.')
    existing = next((x for x in state[kind] if norm(x['name'])==norm(name)),None)
    if existing:
        return existing['id']
    item={'id':uid(),'name':name}
    state[kind].append(item)
    return item['id']


def archive_rows(state, ids):
    rows=[r for r in state['rows'] if r['id'] in ids]
    for row in rows:
        state['trash'].append({'id':row['id'],'deleted_at':now(),'row':row,
             'supplier':deepcopy(entity(state,'suppliers',row['supplier_id'])),
             'category':deepcopy(entity(state,'categories',row['category_id'])),
             'companies':[deepcopy(entity(state,'companies',i)) for i in dict.fromkeys(row['company_ids']+list(row.get('company_prices',{}))+list(entity(state,'suppliers',row['supplier_id']).get('prices',{})))],
             'fields':deepcopy(state['fields'])})
    state['rows']=[r for r in state['rows'] if r['id'] not in ids]


@app.post('/api/action')
def action(body:dict,u=Depends(user)):
    with db() as c:
        c.execute('BEGIN IMMEDIATE')
        u=current_access(c,u['id'])
        state,revision=read_state(c)
        if u['role']!='admin':
            p=body.get('data',{})
            row=next((row for row in state['rows'] if row['id']==p.get('id') and visible_row(row,u)),None)
            if body.get('op')!='row_company_price' or p.get('company_id') not in u['company_ids'] or not row:
                raise HTTPException(403,'U kunt alleen productprijzen van uw eigen bedrijf wijzigen.')
        check_revision(body,revision)
        before=deepcopy(state); op=body.get('op'); p=body.get('data',{}); label='Wijziging opgeslagen'
        if op=='row_save':
            supplier_id=p.get('supplier_id')
            if supplier_id:
                entity(state,'suppliers',supplier_id)
            else:
                name=str(p.get('supplier_name','')).strip()
                matches=[s for s in state['suppliers'] if norm(s['name'])==norm(name)]
                if matches and p.get('duplicate_decision')!='separate':
                    raise HTTPException(422,'Deze leveranciersnaam bestaat al. Kies de bestaande leverancier of bevestig een afzonderlijke leverancier.')
                if not name:
                    raise HTTPException(422,'De leveranciersnaam is verplicht.')
                supplier_id=uid(); state['suppliers'].append({'id':supplier_id,'name':name})
            category_id=p.get('category_id'); entity(state,'categories',category_id)
            if p.get('id'):
                row=entity(state,'rows',p['id'])
            else:
                row={'id':uid(),'values':{}};state['rows'].append(row)
            row.update({'supplier_id':supplier_id,'category_id':category_id,'company_ids':p.get('company_ids',[]),'values':p.get('values',{})})
            update_product(row,p)
            if 'company_prices' in p:update_company_prices(state,row,p['company_prices'])
            label='Productregel opgeslagen'
        elif op=='row_company_price':
            row=entity(state,'rows',p.get('id'))
            if 'price' not in p:raise HTTPException(422,'De bedrijfsprijs ontbreekt.')
            update_company_prices(state,row,{p.get('company_id'):p['price']})
            label='Productprijs voor '+entity(state,'companies',p.get('company_id'))['name']+' opgeslagen'
        elif op in ['row_price','row_product']:
            row=entity(state,'rows',p.get('id'))
            key='price' if op=='row_price' else 'product_name'
            if key not in p:raise HTTPException(422,'De productnaam of prijs ontbreekt.')
            update_product(row,{key:p[key]})
            label='Productprijs opgeslagen' if op=='row_price' else 'Productnaam opgeslagen'
        elif op=='rows_delete':
            ids=p.get('ids',[])
            for i in ids: entity(state,'rows',i)
            archive_rows(state,ids);label=f'{len(ids)} productregel(s) naar prullenbak'
        elif op=='supplier_delete':
            supplier=entity(state,'suppliers',p['id'])
            archive_rows(state,[r['id'] for r in state['rows'] if r['supplier_id']==p['id']])
            state['suppliers']=[s for s in state['suppliers'] if s['id']!=p['id']]
            label=f'Leverancier {supplier["name"]} verwijderd'
        elif op=='supplier_rename':
            supplier=entity(state,'suppliers',p['id']);name=str(p.get('name','')).strip()
            if any(norm(s['name'])==norm(name) and s['id']!=p['id'] for s in state['suppliers']) and not p.get('keep_separate'):
                raise HTTPException(422,'Mogelijke dubbele leverancier. Kies samenvoegen of bevestig afzonderlijk bewaren.')
            supplier['name']=name;label='Leverancier hernoemd; alle productregels bijgewerkt'
        elif op=='supplier_prices':
            supplier=entity(state,'suppliers',p.get('id'))
            prices=p.get('prices')
            if not isinstance(prices,dict):raise HTTPException(422,'Ongeldige totaalprijzen.')
            saved=supplier.setdefault('prices',{})
            for company_id,value in prices.items():
                company=entity(state,'companies',company_id)
                amount=price_amount(value)
                if amount=='':saved.pop(company_id,None)
                else:saved[company_id]=amount
            label=f'Totaalprijzen opgeslagen voor {supplier["name"]}'
        elif op=='supplier_merge':
            source=entity(state,'suppliers',p['source']);target=entity(state,'suppliers',p['target'])
            if source['id']==target['id']: raise HTTPException(422,'Kies twee verschillende leveranciers.')
            for company_id,amount in source.get('prices',{}).items():
                current=target.get('prices',{}).get(company_id)
                if current is not None and price_amount(current)!=price_amount(amount):
                    choice=p.get('price_choices',{}).get(company_id)
                    if choice not in ['source','target']:
                        raise HTTPException(422,f'Beide leveranciers hebben een verschillende totaalprijs voor {entity(state,"companies",company_id)["name"]}. Kies welke prijs behouden blijft; prijzen worden niet opgeteld.')
                    if choice=='target':continue
                target.setdefault('prices',{})[company_id]=amount
            for r in state['rows']:
                if r['supplier_id']==source['id']:r['supplier_id']=target['id']
            state['suppliers']=[s for s in state['suppliers'] if s['id']!=source['id']]
            label=f'{source["name"]} samengevoegd met {target["name"]}; alle productregels behouden'
        elif op=='rows_bulk':
            for i in p.get('ids',[]):
                r=entity(state,'rows',i)
                if p.get('category_id'):r['category_id']=p['category_id']
                for company,value in p.get('companies',{}).items():
                    entity(state,'companies',company)
                    r['company_ids']=[i for i in r['company_ids'] if i!=company]
                    if value:r['company_ids'].append(company)
                r['values'].update(p.get('values',{}))
            label=f'{len(p.get("ids",[]))} productregels gezamenlijk gewijzigd'
        elif op=='entity_save':
            kind=p.get('kind')
            if kind not in ['categories','companies','fields']:raise HTTPException(422,'Onbekend type.')
            item=entity(state,kind,p['id']) if p.get('id') else {'id':uid()}
            item['name']=str(p.get('name','')).strip()
            if kind=='fields':
                item.update({'type':p.get('type','text'),'options':[str(x).strip() for x in p.get('options',[]) if str(x).strip()]})
                if item['type']=='select' and not item['options']:raise HTTPException(422,'Voeg minstens één keuze toe.')
            if not p.get('id'):state[kind].append(item)
            label='Instelling opgeslagen'
        elif op=='entity_delete':
            kind=p.get('kind')
            if kind not in ['categories','companies','fields']:raise HTTPException(422,'Onbekend type.')
            item=entity(state,kind,p['id']);target=p.get('target')
            affected=[r for r in state['rows'] if r['category_id']==item['id']] if kind=='categories' else [r for r in state['rows'] if item['id'] in (r['company_ids'] if kind=='companies' else r['values'])]
            priced=[s for s in state['suppliers'] if item['id'] in s.get('prices',{})] if kind=='companies' else []
            product_priced=[r for r in state['rows'] if item['id'] in r.get('company_prices',{})] if kind=='companies' else []
            access_rows=[a for a in c.execute("SELECT * FROM user_access WHERE role='company'") if item['id'] in json.loads(a['company_ids'])] if kind=='companies' else []
            invites=[a for a in c.execute("SELECT * FROM invitations WHERE role='company' AND used_at IS NULL AND canceled_at IS NULL AND expires>?",(time.time(),)) if item['id'] in json.loads(a['company_ids'])] if kind=='companies' else []
            if (affected or priced or product_priced or access_rows or invites) and not target:
                raise HTTPException(422,'Er zijn bestaande koppelingen. Hernoem om deze te behouden, of wijs ze opnieuw toe vóór verwijderen.')
            if target:
                if target==item['id']:raise HTTPException(422,'Kies een andere bestemming.')
                entity(state,kind,target)
                for access in access_rows:
                    companies=list(dict.fromkeys(target if key==item['id'] else key for key in json.loads(access['company_ids'])))
                    c.execute('UPDATE user_access SET company_ids=?,version=version+1 WHERE user_id=?',(js(companies),access['user_id']))
                    c.execute('DELETE FROM sessions WHERE user_id=?',(access['user_id'],))
                for invitation in invites:
                    companies=list(dict.fromkeys(target if key==item['id'] else key for key in json.loads(invitation['company_ids'])))
                    c.execute('UPDATE invitations SET company_ids=? WHERE id=?',(js(companies),invitation['id']))
                for row in product_priced:
                    prices=row['company_prices']
                    if target in prices and price_amount(prices[target])!=price_amount(prices[item['id']]):
                        raise HTTPException(422,'Deze productregel heeft al een andere prijs voor het doelbedrijf. Pas eerst de prijzen aan of wis één bedrag; er is niets overschreven.')
                    prices[target]=prices.pop(item['id'])
                for supplier in priced:
                    prices=supplier['prices']
                    if target in prices and price_amount(prices[target])!=price_amount(prices[item['id']]):
                        raise HTTPException(422,f'{supplier["name"]} heeft al een andere totaalprijs voor het doelbedrijf. Pas eerst de prijzen aan of wis één bedrag; er is niets overschreven.')
                    prices[target]=prices.pop(item['id'])
                for r in affected:
                    if kind=='categories':r['category_id']=target
                    elif kind=='companies':r['company_ids']=list(dict.fromkeys([target if i==item['id'] else i for i in r['company_ids']]))
                    else:
                        if target in r['values'] and r['values'][target] not in ['',None]:raise HTTPException(422,'Het doelveld bevat al gegevens. Er is niets overschreven.')
                        r['values'][target]=r['values'].pop(item['id'])
            state[kind]=[x for x in state[kind] if x['id']!=item['id']];label='Instelling verwijderd; koppelingen behouden'
        elif op=='trash_restore':
            t=deepcopy(entity(state,'trash',p['id']));r=deepcopy(t['row'])
            # Resolve company IDs in both restored rows and supplier totals before
            # bringing back a deleted supplier. Existing suppliers keep newer totals.
            for company in t['companies']:
                replacement=next((x for x in state['companies'] if x['id']==company['id'] or norm(x['name'])==norm(company['name'])),None)
                if replacement and replacement['id']!=company['id']:
                    r['company_ids']=[replacement['id'] if x==company['id'] else x for x in r['company_ids']]
                    product_prices=r.get('company_prices',{})
                    if company['id'] in product_prices:
                        amount=product_prices.pop(company['id'])
                        if replacement['id'] in product_prices and price_amount(product_prices[replacement['id']])!=price_amount(amount):
                            raise HTTPException(422,'Twee herstelde productprijzen komen op hetzelfde bedrijf uit. Gebruik de wijzigingsgeschiedenis om de volledige toestand te herstellen.')
                        product_prices[replacement['id']]=amount
                    prices=t['supplier'].get('prices',{})
                    if company['id'] in prices:
                        amount=prices.pop(company['id'])
                        if replacement['id'] in prices and prices[replacement['id']]!=amount:raise HTTPException(422,'Twee herstelde totaalprijzen komen op hetzelfde bedrijf uit. Gebruik de wijzigingsgeschiedenis om de volledige toestand te herstellen.')
                        prices[replacement['id']]=amount
            for kind,items in [('suppliers',[t['supplier']]),('categories',[t['category']]),('companies',t['companies']),('fields',t['fields'])]:
                for item in items:
                    if not any(x['id']==item['id'] for x in state[kind]):
                        replacement=next((x for x in state[kind] if kind!='suppliers' and norm(x['name'])==norm(item['name'])),None)
                        if replacement:
                            if kind=='categories':r['category_id']=replacement['id']
                            elif kind=='companies':r['company_ids']=[replacement['id'] if x==item['id'] else x for x in r['company_ids']]
                            elif kind=='fields' and item['id'] in r['values']:r['values'][replacement['id']]=r['values'].pop(item['id'])
                        else:state[kind].append(item)
            r['company_ids']=list(dict.fromkeys(r['company_ids']))
            state['rows'].append(r);state['trash']=[x for x in state['trash'] if x['id']!=p['id']];label='Verwijderde productregel hersteld'
        elif op=='history_restore':
            record=c.execute('SELECT * FROM audit WHERE id=?',(p.get('id'),)).fetchone()
            if not record:raise HTTPException(404,'Wijziging niet gevonden.')
            state=json.loads(record['before_data']);label='Gegevens teruggezet naar vóór: '+record['label']
        else:raise HTTPException(422,'Onbekende bewerking.')
        if 'effective_date' in p and op in ['row_save','row_company_price','row_price','row_product']:
            from backend.price_history import effective_date, record_prices
            day=effective_date(p['effective_date'])
            row['information_date']=day
            if op=='row_company_price':values={p['company_id']:p['price']}
            elif op=='row_price':values={'general':p['price']}
            elif op=='row_save':values={**p.get('company_prices',{}),**({'general':p['price']} if 'price' in p else {})}
            else:values={}
            if values:record_prices(state,before,row,values,day,now(),price_amount)
            label+=' (datum '+day+')'
        result=save(c,before,state,revision,u['email'],label,record_dates=op not in ['history_restore','trash_restore'])
        result['state']=scoped_state(result['state'],u);result['user']=u
        return result


@app.get('/api/history')
def history(u=Depends(admin),offset:int=0):
    with db() as c:
        rows=c.execute('SELECT id,at,actor,label FROM audit ORDER BY at DESC,rowid DESC LIMIT 100 OFFSET ?', (max(0,offset),)).fetchall()
    return [dict(r) for r in rows]


def uploaded_table(content,filename):
    if filename.lower().endswith('.xlsx'):
        try:
            w=load_workbook(io.BytesIO(content),data_only=True,read_only=True)
            s=w['Blad1'] if 'Blad1' in w.sheetnames else w.active
            rows=[[str(x) if x is not None else '' for x in r] for r in s.values]
            # Extra category/field metadata in own exports remains separate from product rows.
            category_names=[str(r[0]).strip() for r in w['Blad3'].values if norm(r[0])] if 'Blad3' in w.sheetnames else []
            field_defs=[]
            if 'Velddefinities' in w.sheetnames:
                for r in list(w['Velddefinities'].values)[1:]:
                    field_defs.append({'name':str(r[0]),'type':str(r[1]),'options':json.loads(r[2] or '[]')})
        except Exception:
            raise HTTPException(422,'Het Excel-bestand kan niet worden gelezen. Gebruik een geldig .xlsx-bestand.')
    elif filename.lower().endswith('.csv'):
        try:
            text=content.decode('utf-8-sig')
            try: dialect=csv.Sniffer().sniff(text[:8192],delimiters=',;\t')
            except csv.Error:dialect=csv.excel
            rows=list(csv.reader(io.StringIO(text),dialect));category_names=[];field_defs=[]
        except Exception:raise HTTPException(422,'Gebruik een CSV-bestand met UTF-8-tekencodering.')
    else:raise HTTPException(422,'Kies een Excel (.xlsx) of CSV (.csv) bestand.')
    first=next((i for i,r in enumerate(rows) if any(norm(x) for x in r)),None)
    if first is None:raise HTTPException(422,'Het bestand is leeg.')
    headers=rows[first]
    size=max(len(headers),max((len(r) for r in rows[first+1:]),default=0))
    headers=headers+['']*(size-len(headers))
    labels=[h.strip() or f'Kolom {i+1}' for i,h in enumerate(headers)]
    data=[{'line':i+first+2,'cells':r+['']*(size-len(r))} for i,r in enumerate(rows[first+1:]) if any(norm(x) for x in r)]
    if len(data)>20000 or size>200:raise HTTPException(422,'Maximaal 20.000 regels en 200 kolommen per import.')
    return {'headers':labels,'rows':data,'categories':category_names,'fields':field_defs,'filename':filename}


@app.post('/api/import/upload')
async def upload(file:UploadFile=File(...),u=Depends(admin)):
    content=await file.read(10*1024*1024+1)
    return import_content(content,file.filename or '',u)


def import_content(content,filename,u):
    if len(content)>10*1024*1024:raise HTTPException(422,'Bestand te groot (maximaal 10 MB).')
    table=uploaded_table(content,filename)
    with db() as c:
        state,_=read_state(c);token=uid()
        c.execute('DELETE FROM imports WHERE created<?',(time.time()-86400,))
        c.execute('INSERT INTO imports VALUES(?,?,?,?)',(token,u['id'],js(table),time.time()))
    mapping={'supplier':None,'category':None,'row_id':None,'supplier_id':None,'product_name':None,'price':None,'companies':{},'fields':{},'prices':{},'company_prices':{}}
    for i,h in enumerate(table['headers']):
        n=norm(h)
        if n in ['leveranciers','leverancier','supplier']:mapping['supplier']=i
        if n in ['wat','productcategorie','categorie','product']:mapping['category']=i
        if n in ['regel-id','row_id']:mapping['row_id']=i
        if n in ['leverancier-id','supplier_id']:mapping['supplier_id']=i
        if n in ['productnaam','productomschrijving']:mapping['product_name']=i
        if n in ['prijs','prijs (€)','productprijs','productprijs (€)']:mapping['price']=i
        for co in state['companies']:
            if n==norm(co['name']):mapping['companies'][co['id']]=i
            if n==norm('Totaalprijs (€) - '+co['name']):mapping['prices'][co['id']]=i
            if n==norm('Prijs (€) - '+co['name']):mapping['company_prices'][co['id']]=i
        for f in state['fields']:
            if n==norm(f['name']):mapping['fields'][f['id']]=i
    return {'token':token,'headers':table['headers'],'sample':table['rows'][:8],'total':len(table['rows']),'mapping':mapping,'new_fields':table['fields']}


def import_table(c,token,u):
    r=c.execute('SELECT * FROM imports WHERE id=? AND user_id=? AND created>?',(token,u['id'],time.time()-86400)).fetchone()
    if not r:raise HTTPException(404,'De importvoorvertoning is verlopen. Kies het bestand opnieuw.')
    return json.loads(r['payload'])


def import_preview(state,table,mapping,mode):
    def cell(r,col):
        try:return r['cells'][int(col)].strip() if col is not None and int(col)>=0 else ''
        except (ValueError,IndexError,TypeError):return ''
    price_only=mapping.get('category') is None and bool(mapping.get('prices'))
    if mapping.get('supplier') is None or (mapping.get('category') is None and not price_only):
        raise HTTPException(422,'Wijs leverancier en productcategorie toe, of leverancier en minstens één totaalprijskolom.')
    previews=[]
    for source in table['rows']:
        name=cell(source,mapping.get('supplier'));category=cell(source,mapping.get('category'))
        supplier_id=cell(source,mapping.get('supplier_id'))
        known=next((s for s in state['suppliers'] if s['id']==supplier_id),None) if supplier_id else None
        exact=[s for s in state['suppliers'] if norm(s['name'])==norm(name)]
        matches=[known] if known else exact
        fuzzy=[s for s in state['suppliers'] if s not in matches and difflib.SequenceMatcher(None,norm(name),norm(s['name'])).ratio()>=0.78]
        candidates=matches+fuzzy
        company_ids=[i for i,col in mapping.get('companies',{}).items() if norm(cell(source,col)) in ['x','1','true','ja']]
        errors=[]
        if not name:errors.append('Leveranciersnaam ontbreekt')
        if not category and not price_only:errors.append('Productcategorie ontbreekt')
        product_name=cell(source,mapping.get('product_name'))
        if len(product_name)>250:errors.append('Productnaam is te lang (maximaal 250 tekens)')
        product_price=None
        price_cell=cell(source,mapping.get('price'))
        if price_cell:
            try:product_price=price_amount(price_cell)
            except HTTPException as error:errors.append(f'Productprijs: {error.detail}')
        company_prices={}
        for company_id,col in mapping.get('company_prices',{}).items():
            company=entity(state,'companies',company_id)
            value=cell(source,col)
            if not value:continue  # Blank import cells preserve existing company prices.
            try:company_prices[company_id]=price_amount(value)
            except HTTPException as error:errors.append(f'Productprijs {company["name"]}: {error.detail}')
        prices={};price_changes=[]
        for company_id,col in mapping.get('prices',{}).items():
            company=entity(state,'companies',company_id)
            value=cell(source,col)
            if not value:continue  # Empty import cells preserve an existing total.
            try:prices[company_id]=price_amount(value)
            except HTTPException as error:errors.append(f'{company["name"]}: {error.detail}');continue
            previous=matches[0].get('prices',{}).get(company_id) if len(matches)==1 else None
            price_changes.append({'company':company['name'],'amount':prices[company_id],'previous':previous,'changed':previous!=prices[company_id]})
        for company,col in mapping.get('companies',{}).items():
            v=norm(cell(source,col))
            if v not in ['','x','1','true','ja','0','false','nee']:errors.append(f'Bedrijfskoppeling "{cell(source,col)}" is ongeldig (gebruik X of leeg)')
        cat=next((x for x in state['categories'] if norm(x['name'])==norm(category)),None)
        matching_rows=[r for r in state['rows'] if r['supplier_id'] in [s['id'] for s in matches] and cat and r['category_id']==cat['id']]
        if product_name:matching_rows=[r for r in matching_rows if norm(r.get('product_name') or category)==norm(product_name)]
        row_id=cell(source,mapping.get('row_id'));by_id=next((r for r in state['rows'] if r['id']==row_id),None)
        target=by_id['id'] if by_id else matching_rows[0]['id'] if len(matching_rows)==1 else None
        preview={'line':source['line'],'name':name,'category':category,'company_ids':company_ids,'prices':prices,'company_prices':company_prices,'price_changes':price_changes,'price_only':price_only,'product_name':product_name,'product_price':product_price,
                 'values':{i:cell(source,col) for i,col in mapping.get('fields',{}).items()},
                 'errors':errors,'candidates':candidates,'exact_ids':[s['id'] for s in matches],
                 'existing_rows':[{'id':r['id'],'name':entity(state,'suppliers',r['supplier_id'])['name'],'category':entity(state,'categories',r['category_id'])['name']} for r in matching_rows],
                 'target':target,'action':'update' if price_only or (mode=='update' and target) else 'add','duplicate_product':bool(matching_rows),
                 'ambiguous':mode=='update' and len(matching_rows)>1 and not by_id}
        previews.append(preview)
    return previews


@app.post('/api/import/preview')
def preview_import(body:dict,u=Depends(admin)):
    with db() as c:
        state,revision=read_state(c);table=import_table(c,body.get('token'),u)
    rows=import_preview(state,table,body.get('mapping',{}),body.get('mode','add'))
    return {'rows':rows,'revision':revision,'total':len(rows),'errors':sum(bool(r['errors']) for r in rows),
            'duplicates':sum(bool(r['candidates'] or r['duplicate_product']) for r in rows)}


@app.post('/api/import/confirm')
def confirm_import(body:dict,u=Depends(admin)):
    with db() as c:
        c.execute('BEGIN IMMEDIATE');u=check_admin(c,u);state,revision=read_state(c);check_revision(body,revision);before=deepcopy(state)
        table=import_table(c,body.get('token'),u)
        mapping=deepcopy(body.get('mapping',{}))
        # Explicitly chosen export field definitions; never silently change existing field types.
        for definition in body.get('create_fields',[]):
            existing=next((f for f in state['fields'] if norm(f['name'])==norm(definition['name'])),None)
            if not existing:
                existing={'id':uid(),'name':definition['name'],'type':definition.get('type','text'),'options':definition.get('options',[])}
                state['fields'].append(existing)
            col=next((i for i,h in enumerate(table['headers']) if norm(h)==norm(definition['name'])),None)
            if col is not None:mapping.setdefault('fields',{})[existing['id']]=col
        previews=import_preview(state,table,mapping,body.get('mode','add'))
        report={'added':0,'updated':0,'skipped':0,'errors':[],'filename':table['filename'],'prices_updated':0,'product_prices_updated':0,'company_product_prices_updated':0}
        decisions=body.get('decisions',{})
        resolved_names={}
        touched=set()
        imported_prices={}
        for r in previews:
            d=decisions.get(str(r['line']),{})
            action=d.get('action',r['action'])
            if action=='skip':report['skipped']+=1;continue
            if r['errors']:report['errors'].extend([f'Regel {r["line"]}: {e}' for e in r['errors']]);continue
            if r['ambiguous'] and not d.get('target') and action=='update':
                report['errors'].append(f'Regel {r["line"]}: meerdere bestaande productregels; kies een bestemming of toevoegen.');continue
            choice=d.get('supplier')
            if r['candidates'] and choice is None:
                report['errors'].append(f'Regel {r["line"]}: beoordeel eerst de mogelijke dubbele leverancier.');continue
            if choice and choice!='new':
                supplier_id=entity(state,'suppliers',choice)['id']
            elif choice=='new':
                key=norm(r['name'])
                if key not in resolved_names:
                    supplier_id=uid();state['suppliers'].append({'id':supplier_id,'name':r['name']});resolved_names[key]=supplier_id
                supplier_id=resolved_names[key]
            else:supplier_id=find_or_create(state,'suppliers',r['name'])
            supplier=entity(state,'suppliers',supplier_id)
            for company_id,amount in r['prices'].items():
                key=(supplier_id,company_id)
                if key in imported_prices and imported_prices[key]!=amount:
                    report['errors'].append(f'Regel {r["line"]}: verschillende totaalprijzen voor dezelfde leverancier en {entity(state,"companies",company_id)["name"]}. Kies één bedrag of sla de tegenstrijdige regel over.');continue
                if key not in imported_prices and supplier.get('prices',{}).get(company_id)!=amount:
                    report['prices_updated']+=1
                imported_prices[key]=amount
                supplier.setdefault('prices',{})[company_id]=amount
            if r['price_only']:
                report['updated']+=1
                continue
            category_id=find_or_create(state,'categories',r['category'])
            product_values={}
            if r['product_name']:product_values['product_name']=r['product_name']
            if r['product_price'] is not None:product_values['price']=r['product_price']
            if action=='update':
                target=d.get('target',r['target'])
                if not target:
                    report['errors'].append(f'Regel {r["line"]}: geen bestaande productregel gekozen.');continue
                if target in touched:
                    report['errors'].append(f'Regel {r["line"]}: dezelfde doelregel zou tweemaal worden bijgewerkt.');continue
                row=entity(state,'rows',target);touched.add(target)
                # Omitted columns preserve existing values and company links.
                unmapped=[i for i in row['company_ids'] if i not in mapping.get('companies',{})]
                row.update({'supplier_id':supplier_id,'category_id':category_id,'company_ids':list(dict.fromkeys(unmapped+r['company_ids']))})
                if 'price' in product_values and row.get('price')!=product_values['price']:report['product_prices_updated']+=1
                update_product(row,product_values)
                report['company_product_prices_updated']+=sum(row.get('company_prices',{}).get(company_id)!=amount for company_id,amount in r['company_prices'].items())
                update_company_prices(state,row,r['company_prices'])
                row['values'].update(r['values']);report['updated']+=1
            else:
                row={'id':uid(),'supplier_id':supplier_id,'category_id':category_id,'company_ids':r['company_ids'],'values':r['values']}
                update_product(row,product_values)
                update_company_prices(state,row,r['company_prices'])
                report['company_product_prices_updated']+=len(r['company_prices'])
                if 'price' in product_values:report['product_prices_updated']+=1
                state['rows'].append(row);report['added']+=1
        if report['errors']:
            # Atomic import: never partially apply an import containing unresolved errors.
            raise HTTPException(422, {'message':'Import niet uitgevoerd. Bestaande gegevens zijn behouden.','errors':report['errors']})
        for name in table.get('categories',[]):find_or_create(state,'categories',name)
        result=save(c,before,state,revision,u['email'],f'Import {table["filename"]}: {report["added"]} toegevoegd, {report["updated"]} bijgewerkt, {report["skipped"]} overgeslagen, {report["prices_updated"]} totaalprijzen gewijzigd')
        result['report']=report
        c.execute('DELETE FROM imports WHERE id=?',(body['token'],))
        return result


def safe_cell(value):
    value='' if value is None else str(value)
    # Prevent formula execution when CSV/XLSX is opened in spreadsheet software.
    return "'"+value if value.lstrip().startswith(('=','+','-','@')) else value


def product_totals_table(rows, suppliers, companies):
    def total(values):
        # A sum may exceed the maximum for one input price. State prices are
        # already validated; do not apply the per-cell limit to derived totals.
        return format(sum((Decimal(str(value)) if value not in [None,''] else Decimal('0') for value in values), Decimal('0')), '.2f')
    grouped={}
    for row in rows:grouped.setdefault(row['supplier_id'],[]).append(row)
    headers=['Leverancier-ID','Leverancier','Productregels']+['Som productprijzen (€) - '+co['name'] for co in companies]+['Totaal bedrijfsprijzen (€)','Som algemene productprijzen (€)']
    content=[]
    for supplier in suppliers:
        products=grouped.get(supplier['id'],[])
        totals=[total(row.get('company_prices',{}).get(co['id']) for row in products) for co in companies]
        content.append([supplier['id'],supplier['name'],len(products)]+totals+[total(totals),total(row.get('price') for row in products)])
    return headers,content


@app.post('/api/export')
def export(body:dict,u=Depends(admin)):
    with db() as c:state,revision=read_state(c)
    fmt=body.get('format','xlsx');ids=body.get('ids')
    rows=state['rows'] if ids is None else [r for r in state['rows'] if r['id'] in ids]
    price_headers=['Totaalprijs (€) - '+co['name'] for co in state['companies']]
    selected_supplier_ids={r['supplier_id'] for r in rows}
    supplier_ids=body.get('supplier_ids')
    suppliers=[s for s in state['suppliers'] if (supplier_ids is None or s['id'] in supplier_ids) and (ids is None or s['id'] in selected_supplier_ids or (body.get('scope')=='product_totals' and supplier_ids is not None))]
    price_rows=[[s['id'],s['name']]+[s.get('prices',{}).get(co['id'],'') for co in state['companies']] for s in suppliers]
    price_only=body.get('scope')=='prices'
    calculated=body.get('scope')=='product_totals'
    headers=['Regel-ID','Leverancier-ID','Leverancier','Productcategorie','Productnaam','Prijs (€)']+[co['name'] for co in state['companies']]+[f['name'] for f in state['fields']]+['Prijs (€) - '+co['name'] for co in state['companies']]+price_headers
    content=[]
    for r in rows:
        content.append([r['id'],r['supplier_id'],entity(state,'suppliers',r['supplier_id'])['name'],entity(state,'categories',r['category_id'])['name'],r.get('product_name',''),r.get('price','')]
                       +['X' if co['id'] in r['company_ids'] else '' for co in state['companies']]+[r['values'].get(f['id'],'') for f in state['fields']]+[r.get('company_prices',{}).get(co['id'],'') for co in state['companies']]
                       +[entity(state,'suppliers',r['supplier_id']).get('prices',{}).get(co['id'],'') for co in state['companies']])
    if price_only:headers=['Leverancier-ID','Leverancier']+price_headers;content=price_rows
    if calculated:
        company_ids=body.get('company_ids')
        companies=[co for co in state['companies'] if company_ids is None or co['id'] in company_ids]
        headers,content=product_totals_table(rows,suppliers,companies)
    filename='TVE-leveranciers-'+datetime.now().strftime('%Y-%m-%d')
    if fmt=='csv':
        buf=io.StringIO();writer=csv.writer(buf,delimiter=';');writer.writerow([safe_cell(h) for h in headers]);writer.writerows([[safe_cell(v) for v in r] for r in content])
        return Response('\ufeff'+buf.getvalue(),media_type='text/csv; charset=utf-8',headers={'Content-Disposition':f'attachment; filename="{filename}.csv"'})
    if fmt!='xlsx':raise HTTPException(422,'Onbekend exportformaat.')
    w=Workbook();s=w.active;s.title='Opgetelde prijzen' if calculated else 'Totaalprijzen' if price_only else 'Leveranciers';s.append([safe_cell(h) for h in headers])
    for r in content:s.append([safe_cell(v) for v in r])
    s.freeze_panes='E2';s.auto_filter.ref=s.dimensions
    for i in range(1,len(headers)+1):s.column_dimensions[__import__('openpyxl').utils.get_column_letter(i)].width=24
    cats=w.create_sheet('Blad3')
    for category in state['categories']:cats.append([safe_cell(category['name'])])
    fields=w.create_sheet('Velddefinities');fields.append(['Naam','Type','Opties'])
    for f in state['fields']:fields.append([safe_cell(f['name']),f['type'],js(f['options'])])
    if not price_only and not calculated:
        totals=w.create_sheet('Totaalprijzen');totals.append(['Leverancier-ID','Leverancier']+price_headers)
        for r in price_rows:totals.append([safe_cell(v) for v in r])
        totals.freeze_panes='C2';totals.auto_filter.ref=totals.dimensions
        calculated_headers,calculated_rows=product_totals_table(rows,suppliers,state['companies'])
        summary=w.create_sheet('Opgetelde prijzen');summary.append([safe_cell(h) for h in calculated_headers])
        for row in calculated_rows:summary.append([safe_cell(value) for value in row])
        summary.freeze_panes='D2';summary.auto_filter.ref=summary.dimensions
    buf=io.BytesIO();w.save(buf)
    return Response(buf.getvalue(),media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':f'attachment; filename="{filename}.xlsx"'})


@app.get('/api/backup')
def backup(u=Depends(admin)):
    with db() as c:
        c.execute('BEGIN IMMEDIATE')
        state,revision=read_state(c)
        history=[dict(r) for r in c.execute('SELECT * FROM audit ORDER BY rowid')]
    data={'format':'tve-backup','version':1,'created_at':now(),'revision':revision,'state':state,'history':history}
    return Response(js(data),media_type='application/json',headers={'Content-Disposition':'attachment; filename="TVE-volledige-backup.json"'})


@app.post('/api/backup/preview')
async def backup_preview(file:UploadFile=File(...),u=Depends(admin)):
    content=await file.read(50*1024*1024+1)
    return backup_content(content,u)


def backup_content(content,u):
    if len(content)>50*1024*1024:raise HTTPException(422,'Back-up te groot (maximaal 50 MB).')
    try:data=json.loads(content)
    except Exception:raise HTTPException(422,'Ongeldig JSON-bestand.')
    if data.get('format')!='tve-backup' or data.get('version')!=1:raise HTTPException(422,'Dit is geen ondersteunde TVE-back-up.')
    validate_state(data.get('state'))
    history=data.get('history',[])
    if not isinstance(history,list):raise HTTPException(422,'Ongeldige wijzigingsgeschiedenis.')
    for item in history:
        if not all(k in item for k in ['id','at','actor','label','before_data','after_data']):raise HTTPException(422,'Onvolledige wijzigingsgeschiedenis.')
        try:
            validate_state(json.loads(item['before_data']));validate_state(json.loads(item['after_data']))
        except (ValueError,TypeError):raise HTTPException(422,'Beschadigde wijzigingsgeschiedenis.')
    token=uid()
    with db() as c:c.execute('INSERT INTO imports VALUES(?,?,?,?)',(token,u['id'],js(data),time.time()))
    s=data['state']
    return {'token':token,'rows':len(s['rows']),'suppliers':len(s['suppliers']),'fields':len(s['fields']),'trash':len(s['trash']),'history':len(history),'created_at':data.get('created_at')}


@app.post('/api/backup/restore')
def backup_restore(body:dict,u=Depends(admin)):
    with db() as c:
        c.execute('BEGIN IMMEDIATE');u=check_admin(c,u);before,revision=read_state(c);check_revision(body,revision)
        data=import_table(c,body.get('token'),u)
        if data.get('format')!='tve-backup':raise HTTPException(422,'Onjuist herstelbestand.')
        result=save(c,before,data['state'],revision,u['email'],'Volledige back-up hersteld',record_dates=False)
        for record in data.get('history',[]):
            c.execute('INSERT OR IGNORE INTO audit(id,at,actor,label,before_data,after_data) VALUES(?,?,?,?,?,?)',tuple(record[k] for k in ['id','at','actor','label','before_data','after_data']))
        c.execute('DELETE FROM imports WHERE id=?',(body['token'],))
        return result


# Persistent chunks keep uploads and downloads below Vercel's request/response limit.
# They also survive a request landing on a different serverless worker.
FILE_CHUNK=384*1024


def delete_transfer(c,token,u):
    c.execute('DELETE FROM imports WHERE user_id=? AND (id=? OR id LIKE ?)',(u['id'],token,token+':%'))


def store_part(c,token,index,content,u):
    c.execute('INSERT INTO imports VALUES(?,?,?,?)',(token+':'+str(index),u['id'],js({'data':base64.b64encode(content).decode()}),time.time()))


@app.post('/api/files/start')
def start_transfer(body:dict,u=Depends(admin)):
    kind=body.get('kind');size=body.get('size');filename=str(body.get('filename',''))[:255]
    if kind not in ['import','backup']:raise HTTPException(422,'Onbekend bestandstype.')
    limit=(10 if kind=='import' else 50)*1024*1024
    if not isinstance(size,int) or not 0<size<=limit:raise HTTPException(422,f'Bestand moet tussen 1 byte en {limit//1024//1024} MB zijn.')
    token=uid();metadata={'transfer':True,'kind':kind,'size':size,'filename':filename,'next':0,'received':0}
    with db() as c:
        c.execute('DELETE FROM imports WHERE created<?',(time.time()-86400,))
        c.execute('INSERT INTO imports VALUES(?,?,?,?)',(token,u['id'],js(metadata),time.time()))
    return {'token':token,'chunk_size':FILE_CHUNK}


@app.post('/api/files/{token}/part')
def upload_part(token:str,body:dict,u=Depends(admin)):
    index=body.get('index');encoded=body.get('data')
    if not isinstance(index,int) or index<0 or not isinstance(encoded,str) or len(encoded)>FILE_CHUNK*4//3:raise HTTPException(422,'Ongeldig bestandsdeel.')
    try:content=base64.b64decode(encoded,validate=True)
    except ValueError:raise HTTPException(422,'Beschadigd bestandsdeel.')
    if not content:raise HTTPException(422,'Leeg bestandsdeel.')
    with db() as c:
        c.execute('BEGIN IMMEDIATE');metadata=import_table(c,token,u)
        if not metadata.get('transfer'):raise HTTPException(422,'Dit bestand kan niet worden aangevuld.')
        if index<metadata['next']:
            existing=import_table(c,token+':'+str(index),u)
            if existing.get('data')==encoded:return {'received':metadata['received']}
            raise HTTPException(409,'Dit bestandsdeel is al anders opgeslagen. Kies het bestand opnieuw.')
        if index!=metadata['next'] or metadata['received']+len(content)>metadata['size']:raise HTTPException(422,'Onjuiste volgorde of bestandsgrootte.')
        store_part(c,token,index,content,u);metadata['next']+=1;metadata['received']+=len(content)
        c.execute('UPDATE imports SET payload=? WHERE id=?',(js(metadata),token))
    return {'received':metadata['received']}


@app.post('/api/files/{token}/finish')
def finish_transfer(token:str,u=Depends(admin)):
    with db() as c:
        metadata=import_table(c,token,u)
        if not metadata.get('transfer') or metadata['received']!=metadata['size']:raise HTTPException(422,'Het bestand is nog niet volledig ontvangen.')
        content=b''.join(base64.b64decode(import_table(c,token+':'+str(i),u)['data']) for i in range(metadata['next']))
    result=import_content(content,metadata['filename'],u) if metadata['kind']=='import' else backup_content(content,u)
    with db() as c:delete_transfer(c,token,u)
    return result


@app.post('/api/files/download')
def start_download(body:dict,u=Depends(admin)):
    kind=body.get('kind')
    if kind=='backup':response=backup(u);filename='TVE-volledige-backup.json'
    elif kind=='export':
        response=export(body,u);filename='TVE-leveranciers.'+body.get('format','xlsx')
    else:raise HTTPException(422,'Onbekend downloadtype.')
    content=response.body;token=uid();parts=math.ceil(len(content)/FILE_CHUNK)
    with db() as c:
        c.execute('DELETE FROM imports WHERE created<?',(time.time()-86400,))
        c.execute('INSERT INTO imports VALUES(?,?,?,?)',(token,u['id'],js({'download':True,'parts':parts,'size':len(content)}),time.time()))
        for i in range(parts):store_part(c,token,i,content[i*FILE_CHUNK:(i+1)*FILE_CHUNK],u)
    return {'token':token,'parts':parts,'filename':filename,'content_type':response.media_type,'size':len(content)}


@app.get('/api/files/{token}/part/{index}')
def download_part(token:str,index:int,u=Depends(admin)):
    with db() as c:
        metadata=import_table(c,token,u)
        if not metadata.get('download') or not 0<=index<metadata['parts']:raise HTTPException(404,'Bestandsdeel niet gevonden.')
        content=base64.b64decode(import_table(c,token+':'+str(index),u)['data'])
    return Response(content,media_type='application/octet-stream',headers={'Content-Disposition':'attachment; filename="part.bin"'})


@app.post('/api/files/{token}/cancel')
def cancel_transfer(token:str,u=Depends(admin)):
    with db() as c:
        metadata=import_table(c,token,u)
        if not metadata.get('transfer') and not metadata.get('download'):raise HTTPException(422,'Dit is geen tijdelijke bestandsoverdracht.')
        delete_transfer(c,token,u)
    return {'cancelled':True}


from backend.reporting import register as register_reporting
register_reporting(app, db, user, admin, current_access, read_state, now, js, check_revision)

DIST=ROOT/'frontend/dist'
if DIST.exists():
    app.mount('/assets',StaticFiles(directory=DIST/'assets'),name='assets')


@app.get('/{path:path}')
def frontend(path:str):
    if path.startswith('api/'):
        raise HTTPException(404,'Niet gevonden.')
    if not (DIST/'index.html').exists():
        raise HTTPException(503,'Bouw de interface met npm run build in frontend.')
    return FileResponse(DIST/'index.html')
