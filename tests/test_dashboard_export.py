"""Excel dashboard visuals must reflect the same selected company and periods."""
import io
import unittest

from openpyxl import Workbook, load_workbook

import test_app as fixture
from backend.dashboard_export import add_dashboard


class DashboardExportTests(unittest.TestCase):
    get = fixture.AppTests.get
    action = fixture.AppTests.action
    setUp = fixture.AppTests.setUp

    def workbook(self, **body):
        response = self.client.post('/api/export', json={'format': 'xlsx', **body})
        self.assertEqual(response.status_code, 200, response.text if response.status_code != 200 else '')
        return load_workbook(io.BytesIO(response.content))

    def test_dashboard_is_first_with_logo_three_charts_and_selected_company(self):
        state = self.get()['state']
        row = state['rows'][0]
        first, second = state['companies'][:2]
        self.action('row_save', {**row, 'company_ids': [first['id'], second['id']],
                                'company_prices': {first['id']: '10.25', second['id']: '29.50'}})
        expected = self.client.get('/api/reporting', params={'company_id': second['id']}).json()
        original = self.get()
        workbook = self.workbook(company_id=second['id'])
        self.assertEqual(workbook.active.title, 'Dashboard')
        dashboard = workbook['Dashboard']
        self.assertEqual(len(dashboard._charts), 3)
        self.assertEqual(len(dashboard._images), 1)
        self.assertIn(second['name'], str(dashboard['B7'].value))
        self.assertEqual(dashboard._charts[0].series[0].graphicalProperties.solidFill.srgbClr, 'F36B21')
        self.assertFalse(dashboard._charts[0].visible_cells_only)
        self.assertEqual(workbook['Periodebedragen']['B2'].value, float(expected['series'][0]['total']))
        self.assertEqual(workbook['Periodebedragen']['C2'].value, float(expected['series'][0]['group_total']))
        self.assertEqual(len(expected['series']), 1)
        self.assertEqual(workbook['Grafiekgegevens'].max_row, len(state['companies']) + 1)
        self.assertIsNone(workbook['Grafiekgegevens']['B3'].value)
        self.assertEqual(self.get(), original)

    def test_filtered_stats_and_company_counts_keep_reporting_scope(self):
        state = self.get()['state']
        a, b = state['rows'][:2]
        company = state['companies'][0]
        self.action('row_save', {**a, 'company_ids': [company['id']], 'company_prices': {company['id']: '10'}})
        self.action('row_save', {**b, 'company_ids': [company['id']], 'company_prices': {company['id']: '20'}})
        workbook = self.workbook(ids=[a['id']], company_id=company['id'])
        self.assertEqual(workbook['Dashboard']['L14'].value, 1)
        self.assertEqual(workbook['Dashboard']['AF14'].value, 1)
        self.assertEqual(workbook['Leveranciers'].max_row, 2)
        self.assertEqual(workbook['Periodebedragen']['B2'].value, 30)
        self.assertIn('Actieve selectie', workbook['Dashboard']['B70'].value)

    def test_annual_baseline_and_quarters_use_distinct_year_lines_and_gaps(self):
        state = self.get()['state']
        selected = state['companies'][0]['id']
        series = [{'key': '2026', 'total': '1234.56', 'group_total': '5000.00',
                   'current': False, 'future': False, 'granularity': 'year'}]
        for year in [2027, 2028]:
            for quarter in range(1, 5):
                value = None if quarter > 1 else '0.00' if year == 2027 else '234.50'
                series.append({'key': f'{year}-Q{quarter}', 'total': value, 'group_total': value,
                               'current': year == 2028 and quarter == 1,
                               'future': year == 2028 and quarter > 1, 'granularity': 'quarter'})
        report = {'selected_company': selected, 'informative_only': False, 'period': '2028-Q1',
                  'coverage': {'expected': 1, 'missing': 0}, 'deadline': '2028-03-25', 'series': series}
        workbook = Workbook()
        workbook.active.title = 'Leveranciers'
        add_dashboard(workbook, state, state['rows'], report)
        output = io.BytesIO()
        workbook.save(output)
        workbook = load_workbook(io.BytesIO(output.getvalue()))
        source = workbook['Grafiekgegevens']
        # 2026 annual, Q1, Q2, Q3, Q4. Missing quarter values stay absent.
        self.assertEqual([source.cell(row, 1).value for row in range(2, 7)], ['2026', 'Q1', 'Q2', 'Q3', 'Q4'])
        self.assertEqual(source['B2'].value, 1234.56)
        self.assertIsNone(source['B3'].value)
        self.assertEqual(source['C3'].value, 0)
        self.assertEqual(source['D3'].value, 234.5)
        self.assertIsNone(source['C4'].value)
        bars = workbook['Dashboard']._charts[0]
        self.assertEqual(len(bars.series), 3)
        self.assertEqual(len(bars._charts[1].series), 2)
        lines = bars._charts[1]
        self.assertEqual(lines.display_blanks, 'gap')
        self.assertEqual([s.graphicalProperties.line.solidFill.srgbClr for s in lines.series], ['268575', 'B03178'])
        self.assertEqual(bars.y_axis.scaling.max, lines.y_axis.scaling.max)
        self.assertNotEqual(bars.x_axis.axId, lines.x_axis.axId)
        self.assertTrue(lines.x_axis.delete)
        self.assertAlmostEqual(source['N3'].value, 2 + 1 / 3.7)

    def test_specialised_exports_keep_their_original_first_sheet(self):
        self.assertEqual(self.workbook(scope='prices').active.title, 'Totaalprijzen')
        self.assertEqual(self.workbook(scope='product_totals').active.title, 'Opgetelde prijzen')
        response = self.client.post('/api/export', json={'format': 'xlsx', 'company_id': 'unknown'})
        self.assertEqual(response.status_code, 422)
