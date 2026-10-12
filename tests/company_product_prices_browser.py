"""Company/product pricing with isolated data, optionally using published UI assets.

Production supplier data and account credentials are never modified.
"""
import argparse
import json
import os
import socket
import subprocess
import tempfile
import time
from pathlib import Path
import httpx
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright, expect

ROOT=Path(__file__).resolve().parent.parent
parser=argparse.ArgumentParser();parser.add_argument('--published',action='store_true');args=parser.parse_args()
with socket.socket() as sock:sock.bind(('127.0.0.1',0));port=sock.getsockname()[1]
LOCAL=f'http://127.0.0.1:{port}'
BASE='https://tve-leveranciers.vercel.app' if args.published else LOCAL
folder=Path(tempfile.mkdtemp(prefix='tve-company-product-prices-'))
env={**os.environ,'TVE_DB_PATH':str(folder/'prices.sqlite')}
for key in ['SUPABASE_DB_URL','DATABASE_URL','VERCEL']:env.pop(key,None)
log=open(folder/'server.log','w')
client=httpx.Client(timeout=45,follow_redirects=True)
def start():
    process=subprocess.Popen([str(ROOT/'.venv/bin/python'),'-m','uvicorn','backend.app:app','--host','127.0.0.1','--port',str(port)],cwd=ROOT,env=env,stdout=log,stderr=log)
    for _ in range(100):
        try:
            if client.get(LOCAL+'/api/health').is_success:return process
        except httpx.HTTPError:pass
        time.sleep(.1)
    raise RuntimeError('Geïsoleerde testserver kon niet starten.')
process=start();checks=[]
try:
    with sync_playwright() as p:
        browser=p.chromium.launch(executable_path='/usr/bin/chromium',headless=True,args=['--no-sandbox'])
        context=browser.new_context(viewport={'width':1440,'height':1000},accept_downloads=True)
        if args.published:
            assert client.get(BASE+'/api/health').json()['storage']=='Supabase / PostgreSQL'
            assert client.get(BASE+'/api/state').status_code==401
            def transport(route):
                request=route.request
                headers={k:v for k,v in request.all_headers().items() if k.lower() not in ['host','content-length','accept-encoding']}
                local=request.url.startswith(BASE+'/api/')
                target=LOCAL+request.url[len(BASE):] if local else request.url
                if local and headers.get('origin')==BASE:headers['origin']=LOCAL
                client.cookies.clear()
                response=client.request(request.method,target,headers=headers,content=request.post_data_buffer)
                rh={k:v for k,v in response.headers.items() if k.lower() not in ['content-encoding','content-length','transfer-encoding']}
                route.fulfill(status=response.status_code,headers=rh,body=response.content)
            context.route(BASE+'/**',transport)
        page=context.new_page();errors=[]
        page.on('pageerror',lambda error:errors.append(str(error)))
        page.goto(BASE+'/#setup='+(folder/'prices.token').read_text().strip())
        page.get_by_label('E-mailadres',exact=True).fill('bedrijfsprijs@example.test')
        page.get_by_label('Wachtwoord (minstens 12 tekens)',exact=True).fill('isolated-company-price-password')
        page.get_by_role('button',name='Account aanmaken',exact=True).click()
        expect(page.get_by_role('heading',name='Grip op uw leveranciers')).to_be_visible()
        page.get_by_role('button',name='Leveranciers',exact=True).click()
        expect(page.get_by_label('Prijzen invoeren voor',exact=True)).to_have_value(page.get_by_label('Prijzen invoeren voor',exact=True).locator('option').filter(has_text='TVE').first.get_attribute('value'))
        for name,tve,forty in [('Product A','12,34','23,45'),('Product B','35,67','45,67')]:
            page.get_by_role('button',name='Productregel toevoegen',exact=True).click()
            if name=='Product A':page.get_by_label('Leveranciersnaam',exact=True).fill('Bedrijfsprijscontrole')
            else:
                page.get_by_role('button',name='Bestaande leverancier',exact=True).click()
                page.get_by_label('Leverancier',exact=True).select_option(label='Bedrijfsprijscontrole')
            page.get_by_label('Productcategorie',exact=True).select_option(label='Folie')
            page.get_by_label('Productnaam',exact=True).fill(name)
            page.get_by_label('Productprijs TVE (€)',exact=True).fill(tve)
            page.get_by_label('Productprijs FORTY FIVE (€)',exact=True).fill(forty)
            page.get_by_role('button',name='Opslaan',exact=True).click()
            expect(page.get_by_role('dialog')).not_to_be_visible()
        page.get_by_label('Zoek leverancier',exact=True).fill('Bedrijfsprijscontrole')
        expect(page.locator('.product-prices-table tbody tr')).to_have_count(2)
        expect(page.get_by_label('Prijs Bedrijfsprijscontrole, Product A, TVE',exact=True)).to_have_value('12,34')
        expect(page.get_by_label('Prijs Bedrijfsprijscontrole, Product B, TVE',exact=True)).to_have_value('35,67')
        page.get_by_label('Prijzen invoeren voor',exact=True).select_option(label='FORTY FIVE')
        a=page.get_by_label('Prijs Bedrijfsprijscontrole, Product A, FORTY FIVE',exact=True)
        expect(a).to_have_value('23,45');a.fill('88,90');a.press('Enter');expect(a).to_have_value('88,90')
        expect(page.get_by_label('Prijs Bedrijfsprijscontrole, Product B, FORTY FIVE',exact=True)).to_have_value('45,67')
        page.get_by_label('Prijzen invoeren voor',exact=True).select_option(label='TVE')
        expect(page.get_by_label('Prijs Bedrijfsprijscontrole, Product A, TVE',exact=True)).to_have_value('12,34')
        checks.append('Elk bedrijf en product heeft een eigen prijs; een gewijzigde cel laat andere prijzen intact')
        page.get_by_role('button',name='Bedrijfsprijscontrole',exact=True).first.click()
        dialog=page.get_by_role('dialog',name='Alle prijzen per leverancier',exact=True)
        expect(dialog.locator('.supplier-price-matrix tbody tr')).to_have_count(2)
        expect(dialog.locator('.supplier-price-matrix thead th')).to_have_count(6)
        expect(dialog.get_by_label('Overzichtprijs Bedrijfsprijscontrole, Product A, TVE',exact=True)).to_have_value('12,34')
        expect(dialog.get_by_label('Overzichtprijs Bedrijfsprijscontrole, Product A, FORTY FIVE',exact=True)).to_have_value('88,90')
        zero=dialog.get_by_label('Overzichtprijs Bedrijfsprijscontrole, Product A, L&R',exact=True)
        zero.fill('0');zero.press('Enter');expect(zero).to_have_value('0,00')
        checks.append('Klikken op leverancier toont alle producten en alle vijf bedrijven samen, met bewerkbare bedragen en nulprijzen')
        expect(dialog.get_by_label('Totaal productprijzen TVE',exact=True)).to_have_text('€ 48,01')
        expect(dialog.get_by_label('Totaal productprijzen FORTY FIVE',exact=True)).to_have_text('€ 134,57')
        expect(dialog.get_by_label('Totaal alle bedrijfsprijzen',exact=True)).to_have_text('€ 182,58')
        b=dialog.get_by_label('Overzichtprijs Bedrijfsprijscontrole, Product B, TVE',exact=True)
        b.fill('');b.press('Enter')
        expect(dialog.get_by_label('Totaal productprijzen TVE',exact=True)).to_have_text('€ 12,34')
        expect(dialog.get_by_label('Totaal alle bedrijfsprijzen',exact=True)).to_have_text('€ 146,91')
        b.fill('35,67');b.press('Enter')
        expect(dialog.get_by_label('Totaal alle bedrijfsprijzen',exact=True)).to_have_text('€ 182,58')
        checks.append('Leveranciersoverzicht telt per bedrijf en over alle bedrijven op; wissen en aanpassen werken direct bij')
        a=dialog.get_by_label('Overzichtprijs Bedrijfsprijscontrole, Product A, TVE',exact=True)
        with httpx.Client(base_url=LOCAL,timeout=15) as other:
            assert other.post('/api/auth/login',json={'email':'bedrijfsprijs@example.test','password':'isolated-company-price-password'}).is_success
            latest=other.get('/api/state').json();target=next(r for r in latest['state']['rows'] if r.get('product_name')=='Product A')
            company=next(c['id'] for c in latest['state']['companies'] if c['name']=='TVE')
            a.fill('99,50')
            assert other.post('/api/action',json={'op':'row_company_price','data':{'id':target['id'],'company_id':company,'price':'77,10'},'revision':latest['revision']}).is_success
            a.press('Enter');expect(dialog.locator('.cell-error')).to_be_visible();expect(a).to_have_value('99,50')
            a.press('Escape');expect(dialog).to_be_visible();expect(a).to_have_value('77,10')
        checks.append('Gelijktijdige bedrijfsprijswijziging geeft conflict en bewaart de nieuwste prijs; Escape annuleert alleen de cel')
        dialog.get_by_role('button',name='Sluiten',exact=True).click()
        page.get_by_role('button',name='Bedrijfsprijscontrole',exact=True).first.click()
        with page.expect_download() as download:dialog.get_by_role('button',name='Excel',exact=True).click()
        export_path=folder/'bedrijfsprijzen.xlsx';download.value.save_as(str(export_path))
        exported=list(load_workbook(export_path,data_only=True).active.values);headers=exported[0]
        name=headers.index('Productnaam');tve=headers.index('Prijs (€) - TVE');forty=headers.index('Prijs (€) - FORTY FIVE')
        assert {(r[name],r[tve],r[forty]) for r in exported[1:]}=={('Product A','77.10','88.90'),('Product B','35.67','45.67')}
        summary=list(load_workbook(export_path,data_only=True)['Opgetelde prijzen'].values)
        assert summary[1][summary[0].index('Som productprijzen (€) - TVE')]=='112.77'
        assert summary[1][summary[0].index('Totaal bedrijfsprijzen (€)')]=='247.34'
        expect(dialog.get_by_label('Totaal alle bedrijfsprijzen',exact=True)).to_have_text('€ 247,34')
        checks.append('Excel van één leverancier bevat alle productprijzen van alle bedrijven')
        page.screenshot(path=str(ROOT/'test-results/bedrijfsprijzen-desktop.png'),full_page=True,animations='disabled')
        page.set_viewport_size({'width':390,'height':844});page.wait_for_timeout(350)
        assert page.evaluate('document.documentElement.scrollWidth<=innerWidth')
        scroller=dialog.locator('.data-table-scroll')
        assert scroller.evaluate('(el)=>el.scrollWidth>el.clientWidth')
        scroller.evaluate('(el)=>el.scrollLeft=500')
        page.screenshot(path=str(ROOT/'test-results/bedrijfsprijzen-mobiel.png'),full_page=True,animations='disabled')
        checks.append('Mobiel overzicht heeft horizontaal schuivende bedrijfskolommen en een vaste productkolom')
        dialog.get_by_role('button',name='Sluiten',exact=True).click()
        process.terminate();process.wait(10);process=start()
        page.reload();expect(page.get_by_role('heading',name='Grip op uw leveranciers')).to_be_visible()
        page.get_by_role('button',name='Menu openen',exact=True).click()
        page.get_by_role('button',name='Leveranciers',exact=True).click()
        page.get_by_role('button',name='Alle prijzen per leverancier',exact=True).click()
        page.get_by_label('Kies een leverancier',exact=True).select_option(label='Bedrijfsprijscontrole')
        expect(page.get_by_label('Overzichtprijs Bedrijfsprijscontrole, Product A, TVE',exact=True)).to_have_value('77,10')
        expect(page.get_by_label('Overzichtprijs Bedrijfsprijscontrole, Product A, FORTY FIVE',exact=True)).to_have_value('88,90')
        expect(page.get_by_label('Overzichtprijs Bedrijfsprijscontrole, Product B, TVE',exact=True)).to_have_value('35,67')
        checks.append('Bedrijfsprijzen en het overzicht per leverancier blijven bestaan na een echte backendherstart')
        page.get_by_role('button',name='Totaalprijzen per bedrijf',exact=True).click()
        page.get_by_label('Zoek leverancier',exact=True).fill('Bedrijfsprijscontrole')
        expect(page.locator('.calculated-prices-table tbody tr')).to_have_count(1)
        expect(page.get_by_label('Opgetelde prijs Bedrijfsprijscontrole voor TVE',exact=True)).to_have_text('€ 112,77')
        expect(page.get_by_label('Totaal bedrijfsprijzen Bedrijfsprijscontrole',exact=True)).to_have_text('€ 247,34')
        with page.expect_download() as download:page.get_by_role('button',name='Prijzen Excel',exact=True).click()
        totals_path=folder/'totalen.xlsx';download.value.save_as(str(totals_path))
        summary=list(load_workbook(totals_path,data_only=True).active.values)
        assert len(summary)==2 and summary[1][summary[0].index('Totaal bedrijfsprijzen (€)')]=='247.34'
        page.get_by_label('Zoek leverancier',exact=True).fill('Product A')
        expect(page.get_by_label('Totaal bedrijfsprijzen Bedrijfsprijscontrole',exact=True)).to_have_text('€ 166,00')
        page.get_by_label('Filter bedrijf',exact=True).select_option(label='TVE')
        expect(page.get_by_label('Totaal bedrijfsprijzen Bedrijfsprijscontrole',exact=True)).to_have_text('€ 77,10')
        expect(page.get_by_label('Opgetelde prijs Bedrijfsprijscontrole voor TVE',exact=True)).to_have_text('€ 77,10')
        assert page.evaluate('document.documentElement.scrollWidth<=innerWidth')
        checks.append('Totaalprijzen toont automatisch berekende leveranciersbedragen; Excel en product-/bedrijffilters gebruiken dezelfde som')
        assert not errors,errors;checks.append('Geen JavaScript-runtimefouten')
        browser.close()
    suffix='-published' if args.published else ''
    (ROOT/f'test-results/bedrijfsprijzen-browser{suffix}.json').write_text(json.dumps({'passed':checks,'frontend':BASE,'authenticated_api':'isolated local SQLite','production_data_changed':False},ensure_ascii=False,indent=2))
    print(json.dumps({'passed':checks},ensure_ascii=False,indent=2))
finally:
    process.terminate();process.wait(10);log.close();client.close()
