"""Meaningful API integration checks against an isolated SQLite database."""
import io
import json
import os
import tempfile
import unittest
import hashlib
from urllib.parse import urlparse
from pathlib import Path
from openpyxl import load_workbook
from fastapi.testclient import TestClient

TEMP=tempfile.TemporaryDirectory()
test_url=os.environ.get('TVE_TEST_POSTGRES_URL')
if test_url and urlparse(test_url).hostname not in ['127.0.0.1','localhost','::1']:
    raise RuntimeError('Deze destructieve testfixture mag uitsluitend een lokale testdatabase gebruiken.')
os.environ.pop('SUPABASE_DB_URL',None)
os.environ.pop('DATABASE_URL',None)
os.environ.pop('VERCEL',None)
if test_url:os.environ['DATABASE_URL']=test_url
os.environ['TVE_DB_PATH']=str(Path(TEMP.name)/'test.sqlite')
from backend.app import app,db,seed,TOKEN_PATH,now
from backend.storage import IS_POSTGRES


class AppTests(unittest.TestCase):
    def setUp(self):
        self.client=TestClient(app)
        with db() as c:
            c.execute('DELETE FROM access_audit')
            c.execute('DELETE FROM sessions');c.execute('DELETE FROM users');c.execute('DELETE FROM audit');c.execute('DELETE FROM imports');c.execute('DELETE FROM rate_limits')
            c.execute('UPDATE store SET revision=1,payload=?',(json.dumps(seed()),))
            c.execute("DELETE FROM settings WHERE key='setup_code_hash'")
            if IS_POSTGRES:c.execute('INSERT INTO settings VALUES(?,?)',('setup_code_hash',hashlib.sha256(b'isolated-test-code').hexdigest()))
        TOKEN_PATH.write_text('isolated-test-code')
        r=self.client.post('/api/auth/setup',json={'email':'test@example.test','password':'test-password-1234','code':'isolated-test-code'})
        self.assertEqual(r.status_code,200,r.text)

    def get(self):
        r=self.client.get('/api/state');self.assertEqual(r.status_code,200);return r.json()

    def action(self,op,data,expected=200,revision=None):
        r=self.client.post('/api/action',json={'op':op,'data':data,'revision':revision if revision is not None else self.get()['revision']})
        self.assertEqual(r.status_code,expected,r.text)
        return r.json()

    def invite(self,companies=None,email='collega@example.test',role='company',replace_id=None):
        companies=companies if companies is not None else [self.get()['state']['companies'][0]['id']]
        response=self.client.post('/api/users/invite',json={'email':email,'display_name':'Collega','role':role,'company_ids':companies,'replace_id':replace_id})
        self.assertEqual(response.status_code,200,response.text)
        return response.json()

    def accept(self,invitation):
        colleague=TestClient(app)
        response=colleague.post('/api/auth/accept-invitation',json={'token':invitation['token'],'password':'personal-colleague-password'})
        self.assertEqual(response.status_code,200,response.text)
        return colleague

    def test_invitation_creates_personal_account_without_changing_owner_data(self):
        original=self.get();invitation=self.invite()
        visitor=TestClient(app)
        preview=visitor.post('/api/auth/invitation',json={'token':invitation['token']})
        self.assertEqual(preview.status_code,200);self.assertEqual(preview.json()['email'],'collega@example.test')
        listing=self.client.get('/api/users').json()
        self.assertNotIn(invitation['token'],json.dumps(listing));self.assertNotIn('token_hash',json.dumps(listing))
        with db() as c:
            record=c.execute('SELECT token_hash FROM invitations WHERE id=?',(invitation['id'],)).fetchone()
            self.assertEqual(record['token_hash'],hashlib.sha256(invitation['token'].encode()).hexdigest())
        colleague=self.accept(invitation);profile=colleague.get('/api/auth/status').json()['user']
        self.assertEqual(profile['email'],'collega@example.test');self.assertEqual(profile['role'],'company')
        self.assertEqual(self.get(),original)
        self.assertEqual(self.client.get('/api/users').json()['invitations'][0]['status'],'accepted')
        self.assertEqual(visitor.post('/api/auth/accept-invitation',json={'token':invitation['token'],'password':'another-personal-password'}).status_code,410)
        self.assertEqual(colleague.post('/api/auth/logout').status_code,200)
        self.assertEqual(colleague.get('/api/state').status_code,401)
        self.assertEqual(colleague.post('/api/auth/login',json={'email':'collega@example.test','password':'personal-colleague-password'}).status_code,200)

    def test_company_account_receives_only_its_rows_companies_and_prices(self):
        s=self.get()['state'];a,b=s['rows'][:2];c,d=[company['id'] for company in s['companies'][:2]]
        self.action('row_save',{**a,'company_ids':[c,d],'price':'777,77','company_prices':{c:'12,34',d:'998,87'}})
        self.action('row_save',{**b,'company_ids':[d]})
        self.action('supplier_prices',{'id':a['supplier_id'],'prices':{c:'23,45',d:'996,65'}})
        colleague=self.accept(self.invite([c]));scoped=colleague.get('/api/state').json()['state']
        self.assertEqual([company['id'] for company in scoped['companies']],[c])
        self.assertNotIn(b['id'],[row['id'] for row in scoped['rows']])
        row=next(row for row in scoped['rows'] if row['id']==a['id'])
        self.assertEqual(row['company_ids'],[c]);self.assertEqual(row['company_prices'],{c:'12.34'})
        self.assertNotIn('price',row);self.assertNotIn('seed',scoped);self.assertEqual(scoped['trash'],[])
        self.assertNotIn('998.87',json.dumps(scoped));self.assertNotIn('996.65',json.dumps(scoped))
        for path in ['/api/users','/api/history','/api/backup']:
            self.assertEqual(colleague.get(path).status_code,403)
        for path,body in [('/api/export',{'format':'xlsx'}),('/api/files/download',{'kind':'backup'}),('/api/files/start',{'kind':'import','filename':'data.csv','size':100}),('/api/users/invite',{'email':'attacker@example.test','role':'admin'}),('/api/import/preview',{}),('/api/import/confirm',{}),('/api/backup/restore',{})]:
            self.assertEqual(colleague.post(path,json=body).status_code,403,path)

    def test_company_account_can_only_patch_own_visible_product_price(self):
        s=self.get()['state'];a,b=s['rows'][:2];c,d=[company['id'] for company in s['companies'][:2]]
        self.action('row_save',{**a,'company_ids':[c,d],'company_prices':{c:'12,34',d:'23,45'}})
        self.action('row_save',{**b,'company_ids':[d]})
        colleague=self.accept(self.invite([c]));original=self.get()
        requests=[('row_save',a),('row_product',{'id':a['id'],'product_name':'Verboden'}),('row_price',{'id':a['id'],'price':'1'}),('rows_delete',{'ids':[a['id']]}),('entity_save',{'kind':'companies','name':'Verboden'}),('supplier_prices',{'id':a['supplier_id'],'prices':{c:'1'}}),('row_company_price',{'id':a['id'],'company_id':d,'price':'1'}),('row_company_price',{'id':b['id'],'company_id':c,'price':'1'})]
        for op,values in requests:
            response=colleague.post('/api/action',json={'op':op,'data':values,'revision':original['revision']})
            self.assertEqual(response.status_code,403,response.text);self.assertEqual(self.get(),original)
        response=colleague.post('/api/action',json={'op':'row_company_price','data':{'id':a['id'],'company_id':c,'price':'34,56','company_prices':{d:'1'}},'revision':original['revision']})
        self.assertEqual(response.status_code,200,response.text)
        self.assertNotIn('23.45',json.dumps(response.json()['state']))
        stored=next(row for row in self.get()['state']['rows'] if row['id']==a['id'])
        self.assertEqual(stored['company_prices'],{c:'34.56',d:'23.45'});self.assertEqual(stored['company_ids'],[c,d])
        again=colleague.post('/api/action',json={'op':'row_company_price','data':{'id':a['id'],'company_id':c,'price':'99'},'revision':original['revision']})
        self.assertEqual(again.status_code,409)

    def test_invitation_cancel_expiry_reissue_and_invalid_password(self):
        first=self.invite();second=self.invite(replace_id=first['id']);visitor=TestClient(app)
        stale=self.client.post('/api/users/invite',json={'email':'collega@example.test','company_ids':[self.get()['state']['companies'][0]['id']],'replace_id':first['id']})
        self.assertEqual(stale.status_code,409)
        self.assertEqual(visitor.post('/api/auth/invitation',json={'token':first['token']}).status_code,410)
        self.assertEqual(visitor.post('/api/auth/accept-invitation',json={'token':second['token'],'password':'short'}).status_code,422)
        self.assertEqual(visitor.post('/api/auth/invitation',json={'token':second['token']}).status_code,200)
        self.assertEqual(self.client.post('/api/users/invitations/'+second['id']+'/cancel').status_code,200)
        self.assertEqual(visitor.post('/api/auth/invitation',json={'token':second['token']}).status_code,410)
        third=self.invite()
        with db() as c:c.execute('UPDATE invitations SET expires=0 WHERE id=?',(third['id'],))
        self.assertEqual(visitor.post('/api/auth/accept-invitation',json={'token':third['token'],'password':'valid-personal-password'}).status_code,410)
        self.assertEqual(visitor.post('/api/auth/invitation',json={'token':'invalid'}).status_code,410)
        fourth=self.invite()
        response=visitor.post('/api/auth/accept-invitation',json={'token':fourth['token'],'password':'valid-personal-password','email':'hijack@example.test','role':'admin'})
        self.assertEqual(response.status_code,200);profile=visitor.get('/api/auth/status').json()['user']
        self.assertEqual(profile['email'],'collega@example.test');self.assertEqual(profile['role'],'company')
        self.assertEqual(self.client.post('/api/users/invite',json={'email':'collega@example.test','role':'admin'}).status_code,409)

    def test_access_changes_revoke_sessions_and_protect_owner_and_versions(self):
        s=self.get()['state'];c,d=[company['id'] for company in s['companies'][:2]]
        colleague=self.accept(self.invite([c]));profile=colleague.get('/api/auth/status').json()['user']
        response=self.client.post('/api/users/'+profile['id']+'/access',json={'role':'company','company_ids':[d],'active':False,'version':profile['version']})
        self.assertEqual(response.status_code,200,response.text);self.assertEqual(colleague.get('/api/state').status_code,401)
        self.assertEqual(colleague.post('/api/auth/login',json={'email':profile['email'],'password':'personal-colleague-password'}).status_code,401)
        response=self.client.post('/api/users/'+profile['id']+'/access',json={'role':'admin','active':True,'version':profile['version']})
        self.assertEqual(response.status_code,409)
        response=self.client.post('/api/users/'+profile['id']+'/access',json={'role':'company','company_ids':[d],'active':True,'version':profile['version']+1})
        self.assertEqual(response.status_code,200)
        self.assertEqual(colleague.post('/api/auth/login',json={'email':profile['email'],'password':'personal-colleague-password'}).status_code,200)
        self.assertEqual(colleague.get('/api/auth/status').json()['user']['company_ids'],[d])
        owner=self.get()['user']
        response=self.client.post('/api/users/'+owner['id']+'/access',json={'role':'company','company_ids':[c],'active':False,'version':owner['version']})
        self.assertEqual(response.status_code,422);self.assertEqual(self.get()['user'],owner)

    def test_admin_invitation_and_company_reassignment_keep_access_consistent(self):
        administrator=self.accept(self.invite(email='beheerder@example.test',role='admin'))
        self.assertEqual(administrator.get('/api/state').json()['state'],self.get()['state'])
        self.assertEqual(administrator.get('/api/users').status_code,200)
        self.action('entity_save',{'kind':'companies','name':'Nieuw A'})
        self.action('entity_save',{'kind':'companies','name':'Nieuw B'})
        s=self.get()['state'];c,d=[next(company['id'] for company in s['companies'] if company['name']==name) for name in ['Nieuw A','Nieuw B']]
        colleague=self.accept(self.invite([c]));pending=self.invite([c],email='noguitgenodigd@example.test')
        self.action('entity_delete',{'kind':'companies','id':c},expected=422)
        self.action('entity_delete',{'kind':'companies','id':c,'target':d})
        self.assertEqual(colleague.get('/api/state').status_code,401)
        colleague.post('/api/auth/login',json={'email':'collega@example.test','password':'personal-colleague-password'})
        self.assertEqual(colleague.get('/api/auth/status').json()['user']['company_ids'],[d])
        response=TestClient(app).post('/api/auth/invitation',json={'token':pending['token']})
        self.assertEqual(response.status_code,200);self.assertEqual(response.json()['companies'],['Nieuw B'])

    def test_invitation_is_accepted_once_across_independent_connections(self):
        from concurrent.futures import ThreadPoolExecutor
        import threading
        invitation=self.invite();barrier=threading.Barrier(2)
        def redeem():
            with TestClient(app) as client:
                barrier.wait()
                return client.post('/api/auth/accept-invitation',json={'token':invitation['token'],'password':'one-personal-password'}).status_code
        with ThreadPoolExecutor(max_workers=2) as pool:
            results=list(pool.map(lambda _:redeem(),range(2)))
        self.assertEqual(sorted(results),[200,410]);self.assertEqual(len(self.client.get('/api/users').json()['users']),2)

    def test_legacy_owner_migration_preserves_password_sessions_and_company_roles(self):
        import subprocess,sys
        owner=self.get()['user'];colleague=self.accept(self.invite());profile=colleague.get('/api/auth/status').json()['user']
        with db() as c:
            password=c.execute('SELECT password FROM users WHERE id=?',(owner['id'],)).fetchone()[0]
            c.execute('DELETE FROM user_access WHERE user_id=?',(owner['id'],))
            c.execute("DELETE FROM settings WHERE key='access_migrated_v1'")
        code="from backend.app import db; import json;\nwith db() as c: print(json.dumps([dict(row) for row in c.execute('SELECT user_id,role FROM user_access')]))"
        migrated=json.loads(subprocess.check_output([sys.executable,'-c',code],env=os.environ))
        roles={row['user_id']:row['role'] for row in migrated}
        self.assertEqual(roles[owner['id']],'admin');self.assertEqual(roles[profile['id']],'company')
        self.assertEqual(self.get()['user']['id'],owner['id'])
        with db() as c:self.assertEqual(c.execute('SELECT password FROM users WHERE id=?',(owner['id'],)).fetchone()[0],password)

    def test_dataset_backup_preserves_accounts_and_refuses_missing_access_companies(self):
        backup=self.client.get('/api/backup')
        self.action('entity_save',{'kind':'companies','name':'Accountbedrijf'})
        company=next(company for company in self.get()['state']['companies'] if company['name']=='Accountbedrijf')
        colleague=self.accept(self.invite([company['id']]))
        original=self.get();profile=colleague.get('/api/auth/status').json()['user']
        preview=self.client.post('/api/backup/preview',files={'file':('backup.json',backup.content)})
        self.assertEqual(preview.status_code,200,preview.text)
        restore=self.client.post('/api/backup/restore',json={'token':preview.json()['token'],'revision':original['revision']})
        self.assertEqual(restore.status_code,422);self.assertEqual(self.get(),original)
        self.assertEqual(colleague.get('/api/auth/status').json()['user'],profile)
        backup=self.client.get('/api/backup')
        preview=self.client.post('/api/backup/preview',files={'file':('backup.json',backup.content)})
        restore=self.client.post('/api/backup/restore',json={'token':preview.json()['token'],'revision':self.get()['revision']})
        self.assertEqual(restore.status_code,200,restore.text)
        self.assertEqual(colleague.get('/api/auth/status').json()['user'],profile)
        self.assertEqual(len(self.client.get('/api/users').json()['users']),2)

    @unittest.skipUnless(IS_POSTGRES,'Alleen voor de PostgreSQL-provider')
    def test_sqlite_migration_to_postgres_keeps_personal_company_roles_and_invites(self):
        import psycopg,subprocess,sys
        from psycopg import sql
        from psycopg.conninfo import make_conninfo
        from scripts.migrate_supabase import migration_sql,snapshot
        source=Path(TEMP.name)/'account-migration.sqlite'
        environment={key:value for key,value in os.environ.items() if key not in ['DATABASE_URL','SUPABASE_DB_URL','VERCEL']}
        environment['TVE_DB_PATH']=str(source)
        code="""from backend.app import app,TOKEN_PATH
from fastapi.testclient import TestClient
owner=TestClient(app)
owner.get('/api/health')
assert owner.post('/api/auth/setup',json={'email':'migration-owner@example.test','password':'personal-owner-password','code':TOKEN_PATH.read_text().strip()}).status_code==200
company=owner.get('/api/state').json()['state']['companies'][0]['id']
invite=owner.post('/api/users/invite',json={'email':'migration-member@example.test','role':'company','company_ids':[company]}).json()
member=TestClient(app)
assert member.post('/api/auth/accept-invitation',json={'token':invite['token'],'password':'personal-member-password'}).status_code==200
assert owner.post('/api/users/invite',json={'email':'migration-pending@example.test','role':'company','company_ids':[company]}).status_code==200
"""
        result=subprocess.run([sys.executable,'-c',code],env=environment,capture_output=True,text=True)
        self.assertEqual(result.returncode,0,result.stderr)
        migration,_=migration_sql(source,'unused-test-opening-code');before=snapshot(source)
        name='tve_access_migration_'+__import__('uuid').uuid4().hex[:12]
        admin_connection=psycopg.connect(test_url,autocommit=True)
        try:
            admin_connection.execute(sql.SQL('CREATE DATABASE {}').format(sql.Identifier(name)))
            with psycopg.connect(make_conninfo(test_url,dbname=name)) as connection:
                connection.execute(migration);connection.execute(migration)
                rows=connection.execute('SELECT user_id,role,company_ids,active,version,display_name FROM tve.user_access').fetchall()
                roles={row[0]:row[1] for row in rows}
                self.assertEqual(roles,{row['user_id']:row['role'] for row in before['user_access']})
                self.assertEqual(sorted(row[2] for row in rows),sorted(row['company_ids'] for row in before['user_access']))
                self.assertEqual(connection.execute('SELECT count(*) FROM tve.invitations').fetchone()[0],2)
                self.assertEqual(connection.execute('SELECT count(*) FROM tve.users').fetchone()[0],2)
        finally:
            admin_connection.execute(sql.SQL('DROP DATABASE IF EXISTS {} WITH (FORCE)').format(sql.Identifier(name)));admin_connection.close()

    def test_original_excel_import_preserves_63_rows_and_links(self):
        s=self.get()['state'];self.assertEqual(len(s['rows']),63);self.assertEqual(len(s['suppliers']),56)
        self.assertEqual(len({x['category_id'] for x in s['rows']}),35);self.assertEqual(len(s['categories']),39)
        self.assertEqual(len(s['companies']),5)
        w=load_workbook('data/startgegevens.xlsx',data_only=True)
        raw=[r for r in list(w['Blad1'].values)[2:] if r[0] and str(r[0]).strip()]
        for row,original in zip(s['rows'],raw):
            expected=[s['companies'][j]['id'] for j,col in enumerate([2,4,6,8,10]) if str(original[col] or '').strip().lower()=='x']
            self.assertEqual(row['company_ids'],expected)
        multiple=[x for x in s['suppliers'] if sum(r['supplier_id']==x['id'] for r in s['rows'])>1]
        self.assertTrue(multiple)
        for supplier in multiple:
            rows=[r for r in s['rows'] if r['supplier_id']==supplier['id']]
            self.assertEqual(len(rows),len({r['id'] for r in rows}))

    def test_create_edit_rename_and_multiple_products(self):
        s=self.get()['state']
        result=self.action('row_save',{'supplier_name':'Nieuwe leverancier','category_id':s['categories'][0]['id'],'company_ids':[s['companies'][0]['id']],'values':{s['fields'][0]['id']:'M. Test'}})
        supplier=next(x for x in result['state']['suppliers'] if x['name']=='Nieuwe leverancier')
        self.action('row_save',{'supplier_id':supplier['id'],'category_id':s['categories'][1]['id'],'company_ids':[s['companies'][1]['id']],'values':{}})
        result=self.action('supplier_rename',{'id':supplier['id'],'name':'Hernoemde leverancier'})
        self.assertEqual(sum(r['supplier_id']==supplier['id'] for r in result['state']['rows']),2)
        self.assertEqual(next(x['name'] for x in result['state']['suppliers'] if x['id']==supplier['id']),'Hernoemde leverancier')
        self.assertEqual(len(result['state']['rows']),65)
        # A fresh connection reads the durable data, not a browser or server memory cache.
        with db() as c:self.assertEqual(len(json.loads(c.execute('SELECT payload FROM store').fetchone()[0])['rows']),65)

    def test_conflicting_sessions_do_not_overwrite(self):
        before=self.get();s=before['state']
        self.action('entity_save',{'kind':'categories','name':'Nieuwe categorie'})
        self.action('entity_save',{'kind':'companies','name':'Verouderde wijziging'},expected=409,revision=before['revision'])
        self.assertFalse(any(c['name']=='Verouderde wijziging' for c in self.get()['state']['companies']))

    def test_categories_companies_and_extra_field_types(self):
        self.action('entity_save',{'kind':'categories','name':'Testcategorie'})
        self.action('entity_save',{'kind':'companies','name':'Testbedrijf'})
        for kind,options in [('number',[]),('date',[]),('select',['A','B'])]:
            self.action('entity_save',{'kind':'fields','name':'Test '+kind,'type':kind,'options':options})
        s=self.get()['state'];row=s['rows'][0].copy();row['values']={}
        for field in s['fields']:
            if field['name'].startswith('Test '):row['values'][field['id']]={'number':'12.5','date':'2026-10-09','select':'A'}[field['type']]
        self.action('row_save',row)
        number=next(f for f in s['fields'] if f['type']=='number');row['values'][number['id']]='geen getal'
        self.action('row_save',row,expected=422)
        # Linked category and company cannot be deleted without explicit reassignment.
        self.action('entity_delete',{'kind':'categories','id':s['rows'][0]['category_id']},expected=422)
        company=next(c for c in s['companies'] if any(c['id'] in r['company_ids'] for r in s['rows']))
        self.action('entity_delete',{'kind':'companies','id':company['id']},expected=422)
        target=next(c for c in s['companies'] if c['name']=='Testbedrijf')
        res=self.action('entity_delete',{'kind':'companies','id':company['id'],'target':target['id']})
        self.assertFalse(any(company['id'] in r['company_ids'] for r in res['state']['rows']))

    def test_delete_restore_history_and_bulk(self):
        s=self.get()['state'];ids=[r['id'] for r in s['rows'][:3]]
        self.action('rows_bulk',{'ids':ids,'companies':{s['companies'][0]['id']:True},'category_id':s['categories'][0]['id']})
        self.assertTrue(all(s['companies'][0]['id'] in r['company_ids'] for r in self.get()['state']['rows'] if r['id'] in ids))
        result=self.action('rows_delete',{'ids':ids});self.assertEqual(len(result['state']['trash']),3)
        for row_id in ids:self.action('trash_restore',{'id':row_id})
        self.assertEqual(len(self.get()['state']['rows']),63)
        h=self.client.get('/api/history').json();self.assertEqual(len(h),5)
        self.action('history_restore',{'id':h[0]['id']})
        self.assertEqual(len(self.get()['state']['rows']),62)

    def test_duplicate_warning_and_explicit_merge_preserves_rows(self):
        s=self.get()['state'];supplier=s['suppliers'][0]
        form={'supplier_name':'  '+supplier['name'].lower()+'  ','category_id':s['categories'][0]['id'],'company_ids':[],'values':{}}
        self.action('row_save',form,expected=422)
        result=self.action('row_save',{**form,'duplicate_decision':'separate'})
        other=result['state']['suppliers'][-1]
        result=self.action('supplier_merge',{'source':other['id'],'target':supplier['id']})
        self.assertEqual(len(result['state']['rows']),64);self.assertEqual(len(result['state']['suppliers']),56)

    def test_chunked_import_requires_complete_file_and_preserves_data_until_confirmation(self):
        before=self.get()
        content=self.client.post('/api/export',json={'format':'csv'}).content
        start=self.client.post('/api/files/start',json={'kind':'import','filename':'export.csv','size':len(content)}).json()
        token=start['token']
        self.assertEqual(self.client.post('/api/files/'+token+'/finish').status_code,422)
        import base64
        encoded=base64.b64encode(content).decode()
        self.assertEqual(self.client.post('/api/files/'+token+'/part',json={'index':1,'data':encoded}).status_code,422)
        request={'index':0,'data':encoded}
        self.assertEqual(self.client.post('/api/files/'+token+'/part',json=request).status_code,200)
        self.assertEqual(self.client.post('/api/files/'+token+'/part',json=request).status_code,200)
        result=self.client.post('/api/files/'+token+'/finish')
        self.assertEqual(result.status_code,200,result.text);self.assertEqual(result.json()['total'],63)
        self.assertEqual(self.get(),before)
        with db() as c:
            self.assertFalse(c.execute('SELECT 1 FROM imports WHERE id=? OR id LIKE ?',(token,token+':%')).fetchone())

    def test_large_backup_download_upload_and_restore_across_independent_connections(self):
        import base64
        from backend.app import uid,js
        before=self.get()['state'];payload=js(before)
        # A realistic history eventually exceeds Vercel's 4.5 MB response limit.
        count=5*1024*1024//(2*len(payload.encode()))+2
        with db() as c:
            for _ in range(count):c.execute('INSERT INTO audit(id,at,actor,label,before_data,after_data) VALUES(?,?,?,?,?,?)',(uid(),now(),'Test','Bewaarde wijziging',payload,payload))
        response=self.client.post('/api/files/download',json={'kind':'backup'})
        self.assertEqual(response.status_code,200,response.text);job=response.json();self.assertGreater(job['size'],5*1024*1024)
        parts=[]
        for i in range(job['parts']):
            response=self.client.get('/api/files/'+job['token']+'/part/'+str(i))
            self.assertEqual(response.status_code,200);self.assertLessEqual(len(response.content),384*1024);parts.append(response.content)
        content=b''.join(parts);self.assertEqual(json.loads(content)['state'],before)
        outsider=TestClient(app)
        self.assertEqual(outsider.get('/api/files/'+job['token']+'/part/0').status_code,401)
        self.assertEqual(self.client.post('/api/files/'+job['token']+'/cancel').status_code,200)
        start=self.client.post('/api/files/start',json={'kind':'backup','filename':'backup.json','size':len(content)}).json()
        for i,offset in enumerate(range(0,len(content),start['chunk_size'])):
            response=self.client.post('/api/files/'+start['token']+'/part',json={'index':i,'data':base64.b64encode(content[offset:offset+start['chunk_size']]).decode()})
            self.assertEqual(response.status_code,200,response.text)
        preview=self.client.post('/api/files/'+start['token']+'/finish');self.assertEqual(preview.status_code,200,preview.text)
        restored=self.client.post('/api/backup/restore',json={'token':preview.json()['token'],'revision':self.get()['revision']})
        self.assertEqual(restored.status_code,200,restored.text);self.assertEqual(restored.json()['state'],before)

    def import_exported(self,fmt,mode='update'):
        response=self.client.post('/api/export',json={'format':fmt});self.assertEqual(response.status_code,200)
        up=self.client.post('/api/import/upload',files={'file':('export.'+fmt,response.content)}).json()
        p=self.client.post('/api/import/preview',json={'token':up['token'],'mapping':up['mapping'],'mode':mode}).json()
        self.assertEqual(p['total'],63);self.assertEqual(p['errors'],0)
        decisions={str(r['line']):{'supplier':r['exact_ids'][0],'action':r['action'],'target':r['target']} for r in p['rows']}
        result=self.client.post('/api/import/confirm',json={'token':up['token'],'mapping':up['mapping'],'mode':mode,'decisions':decisions,'revision':p['revision']})
        self.assertEqual(result.status_code,200,result.text)
        self.assertEqual(result.json()['report']['updated'],63)
        self.assertEqual(len(result.json()['state']['rows']),63)
        return result.json()

    def test_excel_roundtrip_preserves_ids_links_values(self):
        s=self.get()['state'];r=s['rows'][0].copy();r['values']={s['fields'][0]['id']:'Contact met é & komma,;'};self.action('row_save',r)
        before=self.get()['state'];after=self.import_exported('xlsx')['state']
        self.assertEqual(before['rows'],after['rows'])

    def test_csv_roundtrip_and_filtered_export(self):
        before=self.get()['state'];after=self.import_exported('csv')['state'];self.assertEqual(before['rows'],after['rows'])
        response=self.client.post('/api/export',json={'format':'xlsx','ids':[before['rows'][0]['id']]})
        w=load_workbook(io.BytesIO(response.content),data_only=True);self.assertEqual(w.active.max_row,2)

    def test_supplier_totals_are_shared_by_products_and_survive_restart(self):
        s=self.get()['state'];supplier=next(x for x in s['suppliers'] if sum(r['supplier_id']==x['id'] for r in s['rows'])>1)
        before_rows=s['rows'];companies=s['companies']
        result=self.action('supplier_prices',{'id':supplier['id'],'prices':{companies[0]['id']:'€ 1.234,56',companies[1]['id']:'0'}})
        priced=next(x for x in result['state']['suppliers'] if x['id']==supplier['id'])
        self.assertEqual(priced['prices'],{companies[0]['id']:'1234.56',companies[1]['id']:'0.00'})
        self.assertEqual(result['state']['rows'],before_rows)
        self.action('supplier_rename',{'id':supplier['id'],'name':'Hernoemde leverancier met totaalprijs'})
        row=next(r for r in before_rows if r['supplier_id']==supplier['id']).copy();row['category_id']=s['categories'][0]['id']
        self.action('row_save',row)
        self.assertEqual(next(x for x in self.get()['state']['suppliers'] if x['id']==supplier['id'])['prices'],priced['prices'])
        if not IS_POSTGRES:
            import subprocess,sys
            code='from backend.app import db,read_state; import json;\nwith db() as c: print(json.dumps(read_state(c)[0]))'
            restored=json.loads(subprocess.check_output([sys.executable,'-c',code],env=os.environ))
            self.assertEqual(restored,self.get()['state'])
        self.action('supplier_prices',{'id':supplier['id'],'prices':{companies[0]['id']:''}})
        self.assertEqual(next(x for x in self.get()['state']['suppliers'] if x['id']==supplier['id'])['prices'],{companies[1]['id']:'0.00'})

    def test_calculated_totals_export_is_exact_and_never_changes_stored_prices(self):
        s=self.get()['state'];a,b=s['rows'][:2];c,d=s['companies'][:2]
        self.action('row_save',{**a,'product_name':'Product A','price':'999','company_prices':{c['id']:'0,10',d['id']:'12,34'}})
        self.action('row_save',{**b,'supplier_id':a['supplier_id'],'company_prices':{c['id']:'0,20',d['id']:'23,45'}})
        self.action('supplier_prices',{'id':a['supplier_id'],'prices':{c['id']:'888'}})
        original=self.get()
        for fmt in ['xlsx','csv']:
            response=self.client.post('/api/export',json={'format':fmt,'scope':'product_totals','ids':[a['id'],b['id']],'supplier_ids':[a['supplier_id']]})
            self.assertEqual(response.status_code,200,response.text if fmt=='csv' else '')
            if fmt=='xlsx':
                workbook=load_workbook(io.BytesIO(response.content),data_only=True)
                self.assertEqual(workbook.active.title,'Opgetelde prijzen');values=list(workbook.active.values)
            else:
                import csv
                values=list(csv.reader(io.StringIO(response.content.decode('utf-8-sig')),delimiter=';'))
            headers,row=values;self.assertEqual(row[headers.index('Som productprijzen (€) - '+c['name'])],'0.30')
            self.assertEqual(row[headers.index('Som productprijzen (€) - '+d['name'])],'35.79')
            self.assertEqual(row[headers.index('Totaal bedrijfsprijzen (€)')],'36.09')
            self.assertEqual(row[headers.index('Som algemene productprijzen (€)')],'999.00')
            self.assertEqual(self.get(),original)
        export=self.client.post('/api/export',json={'format':'xlsx','ids':[a['id'],b['id']]})
        workbook=load_workbook(io.BytesIO(export.content),data_only=True)
        values=list(workbook['Opgetelde prijzen'].values)
        self.assertEqual(values[1][values[0].index('Totaal bedrijfsprijzen (€)')],'36.09')

    def test_calculated_totals_export_respects_product_and_company_selection(self):
        s=self.get()['state'];a,b=s['rows'][:2];c,d=s['companies'][:2]
        self.action('row_save',{**a,'company_prices':{c['id']:'999999999999,99',d['id']:'1'}})
        self.action('row_save',{**b,'supplier_id':a['supplier_id'],'company_prices':{c['id']:'999999999999,99',d['id']:'2'}})
        for ids,expected in [([a['id']],'999999999999.99'),([a['id'],b['id']],'1999999999999.98')]:
            response=self.client.post('/api/export',json={'format':'xlsx','scope':'product_totals','ids':ids,'company_ids':[c['id']]})
            self.assertEqual(response.status_code,200,response.text if response.status_code!=200 else '')
            values=list(load_workbook(io.BytesIO(response.content),data_only=True).active.values)
            headers,row=values
            self.assertNotIn('Som productprijzen (€) - '+d['name'],headers)
            self.assertEqual(row[headers.index('Totaal bedrijfsprijzen (€)')],expected)

    def test_company_product_prices_are_independent_and_patch_only_one_cell(self):
        s=self.get()['state'];a,b=s['rows'][:2];c,d=[x['id'] for x in s['companies'][:2]]
        self.action('row_save',{**a,'product_name':'Product A','price':'45,50','company_prices':{c:'12,34',d:'23,45'}})
        self.action('row_save',{**b,'supplier_id':a['supplier_id'],'product_name':'Product B','company_prices':{c:'34,56'}})
        self.action('supplier_prices',{'id':a['supplier_id'],'prices':{c:'999'}})
        before=self.get()['state']
        result=self.action('row_company_price',{'id':a['id'],'company_id':c,'price':'1.234,56'})
        rows={r['id']:r for r in result['state']['rows']}
        self.assertEqual(rows[a['id']]['company_prices'],{c:'1234.56',d:'23.45'})
        self.assertEqual(rows[b['id']]['company_prices'],{c:'34.56'})
        self.assertEqual(rows[a['id']]['price'],'45.50')
        self.assertEqual(rows[a['id']]['company_ids'],a['company_ids'])
        self.assertEqual(rows[a['id']]['values'],a['values'])
        self.assertEqual(result['state']['suppliers'],before['suppliers'])
        self.action('row_save',{**a,'company_prices':{d:'0'}})
        self.action('supplier_rename',{'id':a['supplier_id'],'name':'Hernoemd met bedrijfsprijzen'})
        self.action('row_company_price',{'id':a['id'],'company_id':c,'price':''})
        row=next(r for r in self.get()['state']['rows'] if r['id']==a['id'])
        self.assertEqual(row['company_prices'],{d:'0.00'})
        self.assertEqual(row['price'],'45.50')

    def test_company_product_price_validation_and_conflicts_are_atomic(self):
        original=self.get();s=original['state'];row=s['rows'][0];company=s['companies'][0]['id']
        for amount in ['-1','1,234',True,{},'1e2','1000000000000']:
            self.action('row_company_price',{'id':row['id'],'company_id':company,'price':amount},expected=422)
            self.assertEqual(self.get(),original)
        self.action('row_company_price',{'id':row['id'],'company_id':'unknown','price':'1'},expected=422)
        self.action('row_save',{**row,'company_prices':[]},expected=422)
        self.assertEqual(self.get(),original)
        self.action('row_company_price',{'id':row['id'],'company_id':company,'price':'9,95'})
        saved=self.get()
        self.action('row_company_price',{'id':row['id'],'company_id':company,'price':'99'},expected=409,revision=original['revision'])
        self.assertEqual(self.get(),saved)

    def test_company_product_price_excel_csv_roundtrips(self):
        s=self.get()['state'];a,b=s['rows'][:2];c,d=[x['id'] for x in s['companies'][:2]]
        self.action('row_save',{**a,'product_name':'Product A','price':'1,25','company_prices':{c:'10,50',d:'20,75'}})
        self.action('row_save',{**b,'supplier_id':a['supplier_id'],'category_id':a['category_id'],'product_name':'Product B','company_prices':{c:'0'}})
        self.action('entity_save',{'kind':'companies','id':c,'name':'Bedrijf hernoemd'})
        before=self.get()['state']
        for fmt in ['xlsx','csv']:
            after=self.import_exported(fmt)['state'];self.assertEqual(after,before)
        response=self.client.post('/api/export',json={'format':'xlsx','ids':[a['id'],b['id']]})
        exported=list(load_workbook(io.BytesIO(response.content),data_only=True).active.values)
        headers=exported[0];index=headers.index('Prijs (€) - Bedrijf hernoemd')
        self.assertEqual(len(exported),3)
        self.assertEqual([r[index] for r in exported[1:]],['10.50','0.00'])

    def test_company_product_price_import_add_update_blank_and_errors(self):
        s=self.get()['state'];row=s['rows'][0];company=s['companies'][0]
        supplier=next(x for x in s['suppliers'] if x['id']==row['supplier_id'])
        category=next(x for x in s['categories'] if x['id']==row['category_id'])
        self.action('row_company_price',{'id':row['id'],'company_id':company['id'],'price':'15,25'})
        for amount,mode,expected in [('-1','update',422),('','update',200),('17,80','update',200),('0','add',200)]:
            text=f'Regel-ID;Leverancier;Productcategorie;Productnaam;Prijs (€) - {company["name"]}\n{row["id"]};{supplier["name"]};{category["name"]};Importproduct;{amount}\n'
            up=self.client.post('/api/import/upload',files={'file':('bedrijfsprijzen.csv',text.encode())}).json()
            self.assertIn(company['id'],up['mapping']['company_prices'])
            pre=self.client.post('/api/import/preview',json={'token':up['token'],'mapping':up['mapping'],'mode':mode}).json()
            original=self.get()
            res=self.client.post('/api/import/confirm',json={'token':up['token'],'mapping':up['mapping'],'mode':mode,'revision':pre['revision'],'decisions':{str(pre['rows'][0]['line']):{'supplier':supplier['id'],'target':row['id'],'action':mode if mode=='add' else 'update'}}})
            self.assertEqual(res.status_code,expected,res.text)
            if expected==422:self.assertEqual(self.get(),original)
            else:
                result=res.json()
                self.assertEqual(result['report']['company_product_prices_updated'],0 if amount=='' else 1)
                if mode=='add':self.assertEqual(result['state']['rows'][-1]['company_prices'],{company['id']:'0.00'})
        saved=next(r for r in self.get()['state']['rows'] if r['id']==row['id'])
        self.assertEqual(saved['company_prices'],{company['id']:'17.80'})
        self.assertEqual(saved['company_ids'],row['company_ids'])

    def test_company_removal_preserves_product_prices_and_rejects_overwrite(self):
        self.action('entity_save',{'kind':'companies','name':'Prijsbedrijf A'})
        self.action('entity_save',{'kind':'companies','name':'Prijsbedrijf B'})
        s=self.get()['state'];row=s['rows'][0];c,d=[next(x['id'] for x in s['companies'] if x['name']==name) for name in ['Prijsbedrijf A','Prijsbedrijf B']]
        self.action('row_save',{**row,'company_prices':{c:'12,50',d:'13,50'}})
        before=self.get()
        self.action('entity_delete',{'kind':'companies','id':c},expected=422)
        self.action('entity_delete',{'kind':'companies','id':c,'target':d},expected=422)
        self.assertEqual(self.get(),before)
        self.action('row_company_price',{'id':row['id'],'company_id':d,'price':''})
        result=self.action('entity_delete',{'kind':'companies','id':c,'target':d})
        updated=next(r for r in result['state']['rows'] if r['id']==row['id'])
        self.assertEqual(updated['company_prices'],{d:'12.50'})
        self.assertEqual(updated['company_ids'],row['company_ids'])

    def test_company_prices_restore_backup_and_new_process(self):
        self.action('entity_save',{'kind':'companies','name':'Alleen prijs'})
        s=self.get()['state'];row=s['rows'][0];company=next(x for x in s['companies'] if x['name']=='Alleen prijs')
        self.action('row_company_price',{'id':row['id'],'company_id':company['id'],'price':'19,99'})
        before=self.get()['state'];backup=self.client.get('/api/backup')
        self.action('rows_delete',{'ids':[row['id']]})
        self.action('entity_delete',{'kind':'companies','id':company['id']})
        replacement=self.action('entity_save',{'kind':'companies','name':'Alleen prijs'})['state']['companies'][-1]
        result=self.action('trash_restore',{'id':row['id']})
        restored=next(r for r in result['state']['rows'] if r['id']==row['id'])
        self.assertEqual(restored['company_prices'],{replacement['id']:'19.99'})
        self.assertEqual(restored['company_ids'],row['company_ids'])
        pre=self.client.post('/api/backup/preview',files={'file':('backup.json',backup.content)})
        res=self.client.post('/api/backup/restore',json={'token':pre.json()['token'],'revision':self.get()['revision']})
        self.assertEqual(res.status_code,200,res.text);self.assertEqual(res.json()['state'],before)
        if not IS_POSTGRES:
            import subprocess,sys
            code='from backend.app import db,read_state; import json;\nwith db() as c: print(json.dumps(read_state(c)[0]))'
            restored=json.loads(subprocess.check_output([sys.executable,'-c',code],env=os.environ))
            self.assertEqual(restored,before)

    def test_product_prices_are_independent_and_preserve_company_totals(self):
        s=self.get()['state'];a,b=s['rows'][:2];supplier=s['suppliers'][0];company=s['companies'][0]
        self.action('supplier_prices',{'id':supplier['id'],'prices':{company['id']:'999'}})
        self.action('row_save',{**a,'supplier_id':supplier['id'],'product_name':'Folie A','price':'12,34'})
        self.action('row_save',{**b,'supplier_id':supplier['id'],'category_id':a['category_id'],'product_name':'Folie B','price':'25,60'})
        before=self.get()['state']
        result=self.action('row_price',{'id':a['id'],'price':'1.234,56'})
        rows={r['id']:r for r in result['state']['rows']}
        self.assertEqual(rows[a['id']]['price'],'1234.56');self.assertEqual(rows[b['id']]['price'],'25.60')
        self.assertEqual(rows[a['id']]['company_ids'],a['company_ids']);self.assertEqual(rows[a['id']]['values'],a['values'])
        self.assertEqual(result['state']['suppliers'],before['suppliers'])
        self.action('row_product',{'id':a['id'],'product_name':'Folie A gewijzigd'})
        self.action('supplier_rename',{'id':supplier['id'],'name':'Prijsleverancier gewijzigd'})
        self.action('row_save',{**a,'supplier_id':supplier['id']})
        row=next(r for r in self.get()['state']['rows'] if r['id']==a['id'])
        self.assertEqual(row['price'],'1234.56');self.assertEqual(row['product_name'],'Folie A gewijzigd')
        self.action('row_price',{'id':a['id'],'price':'0'})
        self.assertEqual(next(r for r in self.get()['state']['rows'] if r['id']==a['id'])['price'],'0.00')
        self.action('row_price',{'id':a['id'],'price':''})
        self.assertNotIn('price',next(r for r in self.get()['state']['rows'] if r['id']==a['id']))

    def test_product_price_validation_and_stale_edits_leave_data_unchanged(self):
        original=self.get();row=original['state']['rows'][0]
        for value in ['-0.01','1,234','NaN',True,[],'1000000000000']:
            self.action('row_price',{'id':row['id'],'price':value},expected=422)
            self.assertEqual(self.get(),original)
        self.action('row_product',{'id':row['id'],'product_name':'a'*251},expected=422)
        self.action('row_price',{'id':'onbekend','price':'5'},expected=422)
        self.action('row_price',{'id':row['id'],'price':'9,50'})
        self.action('row_price',{'id':row['id'],'price':'20'},revision=original['revision'],expected=409)
        self.assertEqual(next(r for r in self.get()['state']['rows'] if r['id']==row['id'])['price'],'9.50')

    def test_named_products_and_prices_roundtrip_excel_csv(self):
        s=self.get()['state'];a,b=s['rows'][:2]
        self.action('row_save',{**a,'product_name':'Product A','price':'1,25'})
        self.action('row_save',{**b,'supplier_id':a['supplier_id'],'category_id':a['category_id'],'product_name':'Product B','price':'2,75'})
        before=self.get()['state']
        for fmt in ['xlsx','csv']:
            after=self.import_exported(fmt)['state'];self.assertEqual(after,before)
        supplier=next(x for x in before['suppliers'] if x['id']==a['supplier_id'])
        category=next(x for x in before['categories'] if x['id']==a['category_id'])
        text=f'Leverancier;Productcategorie;Productnaam;Prijs (€)\n{supplier["name"]};{category["name"]};Product A;3,15\n{supplier["name"]};{category["name"]};Product B;4,25\n'
        up=self.client.post('/api/import/upload',files={'file':('productprijzen.csv',text.encode())}).json()
        pre=self.client.post('/api/import/preview',json={'token':up['token'],'mapping':up['mapping'],'mode':'update'}).json()
        self.assertEqual([r['target'] for r in pre['rows']],[a['id'],b['id']]);self.assertEqual(pre['errors'],0)
        decisions={str(r['line']):{'supplier':supplier['id'],'action':'update','target':r['target']} for r in pre['rows']}
        res=self.client.post('/api/import/confirm',json={'token':up['token'],'mapping':up['mapping'],'mode':'update','revision':pre['revision'],'decisions':decisions})
        self.assertEqual(res.status_code,200,res.text);self.assertEqual(res.json()['report']['product_prices_updated'],2)
        rows={r['id']:r for r in res.json()['state']['rows']}
        self.assertEqual(rows[a['id']]['price'],'3.15');self.assertEqual(rows[b['id']]['price'],'4.25')

    def test_product_price_import_errors_are_atomic_and_empty_prices_preserve_values(self):
        s=self.get()['state'];row=s['rows'][0];supplier=next(x for x in s['suppliers'] if x['id']==row['supplier_id']);category=next(x for x in s['categories'] if x['id']==row['category_id'])
        self.action('row_price',{'id':row['id'],'price':'15,25'})
        for amount,expected in [('-1',422),('',200)]:
            text=f'Regel-ID;Leverancier;Productcategorie;Prijs (€)\n{row["id"]};{supplier["name"]};{category["name"]};{amount}\n'
            up=self.client.post('/api/import/upload',files={'file':('prijs.csv',text.encode())}).json()
            pre=self.client.post('/api/import/preview',json={'token':up['token'],'mapping':up['mapping'],'mode':'update'}).json()
            original=self.get()
            res=self.client.post('/api/import/confirm',json={'token':up['token'],'mapping':up['mapping'],'mode':'update','revision':pre['revision'],'decisions':{str(pre['rows'][0]['line']):{'supplier':supplier['id'],'target':row['id'],'action':'update'}}})
            self.assertEqual(res.status_code,expected,res.text)
            if expected==422:self.assertEqual(self.get(),original)
            self.assertEqual(next(r for r in self.get()['state']['rows'] if r['id']==row['id'])['price'],'15.25')

    def test_product_prices_survive_delete_restore_backup_and_new_process(self):
        row=self.get()['state']['rows'][0]
        self.action('row_save',{**row,'product_name':'Prijs met herstel','price':'19,99'})
        before=self.get()['state'];backup=self.client.get('/api/backup')
        self.action('rows_delete',{'ids':[row['id']]})
        self.action('trash_restore',{'id':row['id']})
        restored=next(r for r in self.get()['state']['rows'] if r['id']==row['id'])
        self.assertEqual(restored['product_name'],'Prijs met herstel');self.assertEqual(restored['price'],'19.99')
        pre=self.client.post('/api/backup/preview',files={'file':('backup.json',backup.content)})
        res=self.client.post('/api/backup/restore',json={'token':pre.json()['token'],'revision':self.get()['revision']})
        self.assertEqual(res.status_code,200,res.text);self.assertEqual(res.json()['state'],before)
        if not IS_POSTGRES:
            import subprocess,sys
            code='from backend.app import db,read_state; import json;\nwith db() as c: print(json.dumps(read_state(c)[0]))'
            restored=json.loads(subprocess.check_output([sys.executable,'-c',code],env=os.environ))
            self.assertEqual(restored,before)

    def test_supplier_price_validation_and_concurrent_edits(self):
        state=self.get();s=state['state'];supplier=s['suppliers'][0];company=s['companies'][0]
        for value in ['-1','1,234','1e3','NaN',True,{},'1000000000000','1.23,45']:
            self.action('supplier_prices',{'id':supplier['id'],'prices':{company['id']:value}},expected=422)
            self.assertEqual(self.get(),state)
        self.action('supplier_prices',{'id':supplier['id'],'prices':{'onbekend':'12'}},expected=422)
        self.action('supplier_prices',{'id':supplier['id'],'prices':{company['id']:'12,50'}})
        self.action('supplier_prices',{'id':supplier['id'],'prices':{company['id']:'99'}},expected=409,revision=state['revision'])
        self.assertEqual(next(x for x in self.get()['state']['suppliers'] if x['id']==supplier['id'])['prices'][company['id']],'12.50')

    def test_supplier_totals_export_import_and_filtered_prices(self):
        s=self.get()['state'];supplier=next(x for x in s['suppliers'] if sum(r['supplier_id']==x['id'] for r in s['rows'])>1);company=s['companies'][0]
        self.action('supplier_prices',{'id':supplier['id'],'prices':{company['id']:'127,89'}})
        before=self.get()['state']
        for fmt in ['csv','xlsx']:
            self.assertEqual(self.import_exported(fmt)['state']['suppliers'],before['suppliers'])
            export=self.client.post('/api/export',json={'format':fmt,'scope':'prices','supplier_ids':[supplier['id']]})
            up=self.client.post('/api/import/upload',files={'file':('prijzen.'+fmt,export.content)}).json()
            self.assertEqual(up['total'],1);self.assertIsNone(up['mapping']['category']);self.assertEqual(len(up['mapping']['prices']),5)
            pre=self.client.post('/api/import/preview',json={'token':up['token'],'mapping':up['mapping']}).json()
            self.assertTrue(pre['rows'][0]['price_only']);self.assertEqual(pre['errors'],0)
            result=self.client.post('/api/import/confirm',json={'token':up['token'],'mapping':up['mapping'],'decisions':{str(pre['rows'][0]['line']):{'supplier':supplier['id'],'action':'update'}},'revision':pre['revision']})
            self.assertEqual(result.status_code,200,result.text);self.assertEqual(result.json()['state'],before)
        export=self.client.post('/api/export',json={'format':'xlsx'})
        w=load_workbook(io.BytesIO(export.content),data_only=True)
        self.assertEqual(w['Totaalprijzen'].max_row,57)
        matched=[r for r in w['Totaalprijzen'].values if r[0]==supplier['id']]
        self.assertEqual(len(matched),1);self.assertEqual(matched[0][2],'127.89')

    def test_price_import_reviews_changes_and_rejects_conflicting_totals_atomically(self):
        s=self.get()['state'];supplier=s['suppliers'][0];company=s['companies'][0]
        self.action('supplier_prices',{'id':supplier['id'],'prices':{company['id']:'10'}})
        text=f'Leverancier-ID;Leverancier;Totaalprijs (€) - {company["name"]}\n{supplier["id"]};{supplier["name"]};12,50\n{supplier["id"]};{supplier["name"]};13\n'
        up=self.client.post('/api/import/upload',files={'file':('prijzen.csv',text.encode('utf-8'))}).json()
        pre=self.client.post('/api/import/preview',json={'token':up['token'],'mapping':up['mapping']}).json()
        self.assertEqual(pre['rows'][0]['price_changes'][0]['previous'],'10.00')
        self.assertTrue(pre['rows'][0]['price_changes'][0]['changed'])
        original=self.get();decisions={str(r['line']):{'supplier':supplier['id'],'action':'update'} for r in pre['rows']}
        body={'token':up['token'],'mapping':up['mapping'],'decisions':decisions,'revision':pre['revision']}
        result=self.client.post('/api/import/confirm',json=body)
        self.assertEqual(result.status_code,422);self.assertIn('verschillende totaalprijzen',result.text);self.assertEqual(self.get(),original)
        decisions[str(pre['rows'][1]['line'])]['action']='skip'
        result=self.client.post('/api/import/confirm',json=body)
        self.assertEqual(result.status_code,200,result.text);self.assertEqual(result.json()['report']['prices_updated'],1)
        self.assertEqual(next(x for x in result.json()['state']['suppliers'] if x['id']==supplier['id'])['prices'][company['id']],'12.50')

    def test_prices_follow_company_rename_reassignment_and_supplier_merge_choices(self):
        s=self.get()['state'];a,b=s['suppliers'][:2];source,target=s['companies'][:2]
        self.action('supplier_prices',{'id':a['id'],'prices':{source['id']:'10',target['id']:'20'}})
        self.action('entity_save',{'kind':'companies','id':source['id'],'name':'Hernoemd prijsbedrijf'})
        self.action('entity_delete',{'kind':'companies','id':source['id']},expected=422)
        self.action('entity_delete',{'kind':'companies','id':source['id'],'target':target['id']},expected=422)
        self.action('supplier_prices',{'id':a['id'],'prices':{target['id']:''}})
        result=self.action('entity_delete',{'kind':'companies','id':source['id'],'target':target['id']})
        self.assertEqual(next(x for x in result['state']['suppliers'] if x['id']==a['id'])['prices'],{target['id']:'10.00'})
        self.action('supplier_prices',{'id':b['id'],'prices':{target['id']:'30'}})
        original=self.get();self.action('supplier_merge',{'source':a['id'],'target':b['id']},expected=422);self.assertEqual(self.get(),original)
        result=self.action('supplier_merge',{'source':a['id'],'target':b['id'],'price_choices':{target['id']:'source'}})
        self.assertEqual(next(x for x in result['state']['suppliers'] if x['id']==b['id'])['prices'][target['id']],'10.00')
        self.assertEqual(len(result['state']['rows']),63)

    def test_prices_in_backup_deleted_supplier_restore_and_old_backup(self):
        s=self.get()['state'];supplier=s['suppliers'][0];company=s['companies'][0]
        self.action('supplier_prices',{'id':supplier['id'],'prices':{company['id']:'145,98'}})
        before=self.get()['state'];backup=self.client.get('/api/backup')
        self.action('supplier_delete',{'id':supplier['id']})
        trash=next(t for t in self.get()['state']['trash'] if t['supplier']['id']==supplier['id'])
        restored=self.action('trash_restore',{'id':trash['id']})
        self.assertEqual(next(x for x in restored['state']['suppliers'] if x['id']==supplier['id'])['prices'][company['id']],'145.98')
        pre=self.client.post('/api/backup/preview',files={'file':('backup.json',backup.content)})
        result=self.client.post('/api/backup/restore',json={'token':pre.json()['token'],'revision':self.get()['revision']})
        self.assertEqual(result.status_code,200,result.text);self.assertEqual(result.json()['state'],before)
        old=json.loads(backup.content)
        for x in old['state']['suppliers']:x.pop('prices',None)
        pre=self.client.post('/api/backup/preview',files={'file':('oude-backup.json',json.dumps(old).encode())})
        self.assertEqual(pre.status_code,200,pre.text)

    def test_import_preview_and_atomic_failure(self):
        original=self.get();response=self.client.post('/api/import/upload',files={'file':('fout.csv','Leverancier;Productcategorie;TVE\nVoorbeeld;Folie;ongeldig\nGoed;Textiel;X\n'.encode())})
        up=response.json();p=self.client.post('/api/import/preview',json={'token':up['token'],'mapping':up['mapping']}).json()
        self.assertEqual(p['errors'],1);self.assertEqual(self.get()['revision'],original['revision'])
        result=self.client.post('/api/import/confirm',json={'token':up['token'],'mapping':up['mapping'],'revision':p['revision']})
        self.assertEqual(result.status_code,422);self.assertEqual(self.get(),original)

    def test_backup_preserves_history_extra_fields_and_trash(self):
        s=self.get()['state'];self.action('rows_delete',{'ids':[s['rows'][0]['id']]})
        backup=self.client.get('/api/backup');self.assertEqual(backup.status_code,200)
        before=self.get()['state'];self.action('entity_save',{'kind':'companies','name':'Na back-up'})
        pre=self.client.post('/api/backup/preview',files={'file':('backup.json',backup.content)})
        self.assertEqual(pre.status_code,200,pre.text)
        res=self.client.post('/api/backup/restore',json={'token':pre.json()['token'],'revision':self.get()['revision']})
        self.assertEqual(res.status_code,200,res.text);self.assertEqual(res.json()['state'],before)
        self.assertGreaterEqual(len(self.client.get('/api/history').json()),3)

    def test_auth_protects_all_data_and_csrf(self):
        anon=TestClient(app)
        for path in ['/api/state','/api/history','/api/backup']:
            self.assertEqual(anon.get(path).status_code,401)
        self.assertEqual(anon.post('/api/export',json={'format':'csv'}).status_code,401)
        self.assertEqual(self.client.post('/api/action',json={},headers={'origin':'https://evil.example'}).status_code,403)
        self.client.post('/api/auth/logout',json={});self.assertEqual(self.client.get('/api/state').status_code,401)
        self.assertEqual(self.client.post('/api/auth/login',json={'email':'test@example.test','password':'wrong'}).status_code,401)
        for _ in range(9):self.client.post('/api/auth/login',json={'email':'test@example.test','password':'wrong'})
        self.assertEqual(self.client.post('/api/auth/login',json={'email':'test@example.test','password':'wrong'}).status_code,429)

    @unittest.skipUnless(IS_POSTGRES,'Alleen voor de PostgreSQL-provider')
    def test_postgres_two_connections_serialize_and_reject_stale_revision(self):
        from concurrent.futures import ThreadPoolExecutor
        import threading
        revision=self.get()['revision'];barrier=threading.Barrier(2)
        def update(name):
            client=TestClient(app);client.cookies.update(self.client.cookies);barrier.wait()
            return client.post('/api/action',json={'op':'entity_save','data':{'kind':'categories','name':name},'revision':revision}).status_code
        with ThreadPoolExecutor(max_workers=2) as pool:
            a=pool.submit(update,'Gelijktijdig A');b=pool.submit(update,'Gelijktijdig B')
            self.assertEqual(sorted([a.result(),b.result()]),[200,409])
        self.assertEqual(self.get()['revision'],revision+1)
        self.assertEqual(sum(x['name'].startswith('Gelijktijdig') for x in self.get()['state']['categories']),1)

    @unittest.skipUnless(IS_POSTGRES,'Alleen voor de PostgreSQL-provider')
    def test_postgres_startup_as_schema_owner_without_database_create(self):
        # Regression for managed PostgreSQL: CREATE SCHEMA IF NOT EXISTS still
        # requires CREATE on the database even when the owned schema exists.
        import psycopg
        from psycopg import sql
        import secrets
        import subprocess
        import sys
        from urllib.parse import urlunsplit,quote
        name='tve_start_'+secrets.token_hex(6);role=name+'_app'
        password=secrets.token_urlsafe(24)
        parsed=urlparse(test_url)
        admin=psycopg.connect(test_url,autocommit=True)
        try:
            admin.execute(sql.SQL('CREATE ROLE {} LOGIN NOSUPERUSER NOCREATEDB NOCREATEROLE NOINHERIT PASSWORD {}').format(sql.Identifier(role),sql.Literal(password)))
            admin.execute(sql.SQL('CREATE DATABASE {}').format(sql.Identifier(name)))
            from psycopg.conninfo import make_conninfo
            with psycopg.connect(make_conninfo(test_url,dbname=name)) as connection:
                connection.execute(sql.SQL('CREATE SCHEMA tve AUTHORIZATION {}').format(sql.Identifier(role)))
            target=urlunsplit(('postgresql',quote(role,safe='')+':'+quote(password,safe='')+'@'+parsed.hostname+':'+str(parsed.port or 5432),'/'+name,'sslmode=disable',''))
            environment={**os.environ,'DATABASE_URL':target}
            environment.pop('SUPABASE_DB_URL',None)
            code="""import hashlib,json
from backend.app import db,health
assert health()['ok']
with db() as c:
 assert not c.execute("SELECT has_database_privilege(current_user,current_database(),'CREATE') AS value").fetchone()['value']
 row=c.execute('SELECT payload,revision FROM store').fetchone()
 state=json.loads(row['payload'])
 assert len(state['rows'])==63 and len(state['suppliers'])==56
 print(json.dumps({'revision':row['revision'],'sha':hashlib.sha256(row['payload'].encode()).hexdigest()}))
"""
            results=[]
            for _ in range(2):
                result=subprocess.run([sys.executable,'-c',code],cwd=Path(__file__).resolve().parent.parent,env=environment,capture_output=True,text=True)
                self.assertEqual(result.returncode,0,result.stderr)
                results.append(json.loads(result.stdout))
            self.assertEqual(results[0],results[1])
        finally:
            admin.execute(sql.SQL('DROP DATABASE IF EXISTS {} WITH (FORCE)').format(sql.Identifier(name)))
            admin.execute(sql.SQL('DROP ROLE IF EXISTS {}').format(sql.Identifier(role)))
            admin.close()

    @unittest.skipUnless(IS_POSTGRES,'Alleen voor de PostgreSQL-provider')
    def test_postgres_private_schema_is_not_readable_as_anon(self):
        import psycopg
        with db() as c:
            c.execute("DO $$ BEGIN IF NOT EXISTS(SELECT 1 FROM pg_roles WHERE rolname='anon') THEN CREATE ROLE anon; END IF; END $$")
        with self.assertRaises(psycopg.errors.InsufficientPrivilege):
            with db() as c:
                c.execute('SET LOCAL ROLE anon')
                c.execute('SELECT payload FROM tve.store')


if __name__=='__main__':unittest.main(verbosity=2)
